from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.drawing.image import Image as WorkbookImage
from PIL import Image


class BoxLabelWorkbookRenderer:
    def __init__(self, template_path: Path) -> None:
        self.template_path = template_path

    def render(self, values: dict[str, str], image: bytes | None) -> bytes:
        workbook = load_workbook(self.template_path)
        sheet = workbook.active
        assert sheet is not None
        for start in (1, 11):
            for offset, key, label in (
                (0, "factory", "工厂"),
                (1, "itemNo", "货号"),
                (2, "productName", "品名"),
                (3, "color", "颜色"),
            ):
                sheet.cell(start + offset, 2, f"{label}：{values[key]}")
        if image is not None:
            if len(image) > 5 * 1024 * 1024:
                raise ValueError("product image is too large")
            with Image.open(BytesIO(image)) as source:
                source.verify()
                width, height = source.size
            if width <= 0 or height <= 0 or width > 6000 or height > 6000:
                raise ValueError("invalid product image size")
            scale = min(142 / width, 189 / height)
            for anchor in ("A2", "A12"):
                picture = WorkbookImage(BytesIO(image))
                picture.width = width * scale
                picture.height = height * scale
                sheet.add_image(picture, anchor)
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()
