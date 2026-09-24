import hashlib
import hmac
from dataclasses import dataclass
from datetime import datetime
from io import BytesIO
from pathlib import PurePosixPath
from xml.etree import ElementTree
from zipfile import BadZipFile, ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.worksheet import Worksheet
from PIL import Image as PillowImage

from app.settings.config import Settings

HEADERS = ("名称", "规格", "数量", "单价", "总金额", "采购单号", "入库单号", "图片")
WIDTHS = (26, 14, 8, 8, 10, 20, 16, 36)
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@dataclass(frozen=True)
class IncomingWorkbookLimits:
    max_source_bytes: int = 20 * 1024 * 1024
    max_zip_entries: int = 20_000
    max_uncompressed_bytes: int = 100 * 1024 * 1024
    max_worksheets: int = 10
    max_data_rows: int = 5_000


class IncomingWorkbookValidationError(ValueError):
    def __init__(self, issues: list[dict[str, str | int]]) -> None:
        self.issues = issues
        super().__init__("来货出入核对表校验失败")


class IncomingWorkbookCodec:
    def __init__(self, settings: Settings, *, limits: IncomingWorkbookLimits | None = None):
        self._secret = settings.incoming_diff_workbook_signing_secret.encode()
        if not self._secret or (
            settings.app_env in {"shared_test", "production"}
            and (
                len(self._secret) < 32
                or not settings._is_real_value(settings.incoming_diff_workbook_signing_secret)
                or "FAKE" in settings.incoming_diff_workbook_signing_secret.upper()
            )
        ):
            raise ValueError("来货出入核对表签名密钥未配置")
        self._limits = limits or IncomingWorkbookLimits()

    def generate(
        self,
        *,
        batch_id: str,
        version: int,
        lines: list[dict[str, object]],
        images: dict[str, bytes],
        generated_at: datetime,
    ) -> tuple[bytes, str, list[dict[str, object]]]:
        if not lines or len(lines) > self._limits.max_data_rows:
            raise ValueError("来货出入明细数量无效")
        workbook = Workbook()
        first_sheet = workbook.active
        assert first_sheet is not None
        workbook.remove(first_sheet)
        sheets: dict[str, Worksheet] = {}
        resized_images: dict[str, bytes] = {}
        snapshot: list[dict[str, object]] = []
        for sequence, raw in enumerate(lines, 1):
            image_id = str(raw["imageId"])
            if image_id not in images:
                raise ValueError("明细来源图片不在本批次")
            factory = str(raw.get("factoryName") or "待确认")
            if factory == "_核对标识" or len(factory) > 31 or any(
                mark in factory for mark in "[]:*?/\\"
            ):
                raise ValueError("工厂名不能用作工作表名称")
            if factory not in sheets:
                if len(sheets) + 2 > self._limits.max_worksheets:
                    raise ValueError("来货出入工厂工作表数量超过上限")
                sheet = workbook.create_sheet(factory)
                sheets[factory] = sheet
                sheet.append((*HEADERS, "_line_token"))
                sheet.row_dimensions[1].height = 24
                for column, width in enumerate(WIDTHS, 1):
                    sheet.column_dimensions[chr(64 + column)].width = width
                sheet.column_dimensions["I"].width = 24
                sheet.column_dimensions["I"].hidden = True
            sheet = sheets[factory]
            token = self._token(batch_id, image_id, sequence)
            row = sheet.max_row + 1
            quantity = raw["quantity"]
            if not isinstance(quantity, int) or isinstance(quantity, bool) or quantity == 0:
                raise ValueError("来货出入数量必须是非零整数")
            values = (
                str(raw.get("productName") or ""), str(raw.get("spec") or ""),
                quantity, None, None,
                str(raw["purchaseOrderId"]) if raw.get("purchaseOrderId") else None,
                None, None, token,
            )
            sheet.append(values)
            for column in (1, 2, 6):
                sheet.cell(row, column).data_type = "s"
            sheet.row_dimensions[row].height = 96
            if image_id not in resized_images:
                resized_images[image_id] = self._resized_image(images[image_id])
            picture = ExcelImage(BytesIO(resized_images[image_id]))
            picture.anchor = f"H{row}"
            sheet.add_image(picture)
            snapshot.append({
                **{key: value for key, value in raw.items() if key != "sourceBusinessDate"},
                "imageId": image_id, "lineToken": token,
                "sheetName": factory, "rowNumber": row,
            })
        signature = self._signature(batch_id, version, [str(x["lineToken"]) for x in snapshot])
        self._metadata(workbook, batch_id, version, generated_at, signature)
        output = BytesIO()
        workbook.save(output)
        content = output.getvalue()
        if len(content) > self._limits.max_source_bytes:
            raise ValueError("生成的核对表超过文件大小上限")
        return content, signature, snapshot

    def parse(
        self,
        content: bytes,
        *,
        batch_id: str,
        version: int,
        signature: str,
        previous_lines: list[dict[str, object]],
    ) -> list[dict[str, object]]:
        workbook = self._open(content)
        if "_核对标识" not in workbook:
            self._fail("invalid_signature", "核对表缺少技术标识")
        metadata = workbook["_核对标识"]
        if metadata.sheet_state != "hidden":
            self._fail("invalid_signature", "核对表技术标识无效")
        if metadata["B1"].value != batch_id:
            self._fail("batch_mismatch", "核对表批次不匹配")
        if metadata["B2"].value != version:
            self._fail("stale_version", "核对表版本已过期")
        expected = {str(line["lineToken"]): line for line in previous_lines}
        if len(expected) != len(previous_lines) or any(
            not str(line["lineToken"]).startswith(f"{line['imageId']}:")
            for line in previous_lines
        ):
            self._fail("invalid_signature", "已存核对表技术标识无效")
        actual_signature = self._signature(batch_id, version, list(expected))
        if not hmac.compare_digest(signature, actual_signature) or not hmac.compare_digest(
            str(metadata["B4"].value), actual_signature
        ):
            self._fail("invalid_signature", "核对表签名无效")
        issues: list[dict[str, str | int]] = []
        parsed: list[dict[str, object]] = []
        seen: set[str] = set()
        sheets = [sheet for sheet in workbook.worksheets if sheet.title != "_核对标识"]
        if not sheets:
            self._fail("invalid_workbook", "核对表缺少业务工作表")
        row_total = 0
        for sheet in sheets:
            if sheet.sheet_state != "visible":
                issues.append(self._issue(
                    "hidden_business_sheet", sheet.title, 1, "业务工作表必须可见"
                ))
            if sheet.max_row < 1 or sheet.max_row - 1 + row_total > self._limits.max_data_rows:
                self._fail("too_many_data_rows", "核对表数据行数超过上限", sheet.title)
            row_total += sheet.max_row - 1
            headers = tuple(sheet.cell(1, index).value for index in range(1, 10))
            if headers != (*HEADERS, "_line_token"):
                issues.append(self._issue("invalid_header", sheet.title, 1, "核对表列名或顺序无效"))
                continue
            if not sheet.column_dimensions["I"].hidden:
                issues.append(self._issue(
                    "invalid_token_column", sheet.title, 1, "行标识列必须隐藏"
                ))
            for row in range(2, sheet.max_row + 1):
                cells = [sheet.cell(row, column) for column in range(1, 10)]
                values = [cell.value for cell in cells]
                if not any(value is not None for value in values):
                    continue
                token = values[8]
                if not isinstance(token, str) or not token:
                    issues.append(self._issue(
                        "missing_token", sheet.title, row, "数据行缺少行标识"
                    ))
                    continue
                if token not in expected:
                    issues.append(self._issue(
                        "foreign_token", sheet.title, row, "行标识不属于本批次"
                    ))
                    continue
                if token in seen:
                    issues.append(self._issue("duplicate_token", sheet.title, row, "行标识重复"))
                    continue
                seen.add(token)
                if any(cell.data_type in {"f", "e"} for cell in cells[:7]):
                    issues.append(self._issue(
                        "invalid_cell", sheet.title, row, "业务列不能使用公式或错误值"
                    ))
                    continue
                quantity = values[2]
                if not isinstance(quantity, int) or isinstance(quantity, bool) or quantity == 0:
                    issues.append(self._issue(
                        "invalid_quantity", sheet.title, row, "数量必须是非零整数"
                    ))
                    continue
                if not all(isinstance(value, str) and value.strip() for value in values[:2]):
                    fields = "、".join(
                        name for name, value in zip(("名称", "规格"), values[:2], strict=True)
                        if not isinstance(value, str) or not value.strip()
                    )
                    issues.append(self._issue(
                        "required_field", sheet.title, row, "名称和规格不能为空"
                    ) | {"field": fields})
                    continue
                source = expected[token]
                parsed.append({
                    **{key: value for key, value in source.items()
                       if key != "sourceBusinessDate"},
                    "sheetName": sheet.title, "rowNumber": row,
                    "factoryName": None if sheet.title == "待确认" else sheet.title,
                    "productName": values[0], "spec": values[1], "quantity": quantity,
                    "unitPrice": str(values[3]) if values[3] is not None else None,
                    "totalAmount": str(values[4]) if values[4] is not None else None,
                    "purchaseOrderId": str(values[5]) if values[5] is not None else None,
                    "inboundOrderNo": str(values[6]) if values[6] is not None else None,
                })
        if issues:
            raise IncomingWorkbookValidationError(issues)
        return parsed

    def advance(
        self, content: bytes, *, batch_id: str, version: int,
        lines: list[dict[str, object]], generated_at: datetime,
    ) -> tuple[bytes, str]:
        workbook = self._open(content)
        signature = self._signature(batch_id, version, [str(line["lineToken"]) for line in lines])
        metadata = workbook["_核对标识"]
        metadata["B2"] = version
        metadata["B3"] = generated_at.isoformat()
        metadata["B4"] = signature
        output = BytesIO()
        workbook.save(output)
        return output.getvalue(), signature

    def _open(self, content: bytes) -> Workbook:
        if len(content) > self._limits.max_source_bytes:
            self._fail("file_too_large", "核对表文件大小超过上限")
        try:
            with ZipFile(BytesIO(content)) as archive:
                entries = archive.infolist()
                if len(entries) > self._limits.max_zip_entries:
                    self._fail("too_many_zip_entries", "核对表内部文件数超过上限")
                if any(
                    PurePosixPath(entry.filename).is_absolute()
                    or ".." in PurePosixPath(entry.filename).parts
                    or "\\" in entry.filename
                    for entry in entries
                ):
                    self._fail("unsafe_archive_path", "核对表包含不安全路径")
                if sum(entry.file_size for entry in entries) > self._limits.max_uncompressed_bytes:
                    self._fail("uncompressed_content_too_large", "核对表解压后超过上限")
                if sum(
                    entry.filename.startswith("xl/worksheets/") and entry.filename.endswith(".xml")
                    for entry in entries
                ) > self._limits.max_worksheets:
                    self._fail("too_many_worksheets", "核对表工作表数超过上限")
                for entry in entries:
                    if entry.filename.endswith(".rels") and any(
                        relation.attrib.get("TargetMode") == "External"
                        for relation in ElementTree.fromstring(archive.read(entry))
                    ):
                        self._fail("external_relationship", "核对表不能引用外部资源")
            return load_workbook(BytesIO(content), data_only=False, read_only=False)
        except (BadZipFile, InvalidFileException, KeyError, ElementTree.ParseError) as error:
            raise IncomingWorkbookValidationError([
                {"code": "invalid_workbook", "message": "核对表文件损坏或格式错误"}
            ]) from error

    @staticmethod
    def _metadata(
        workbook: Workbook, batch_id: str, version: int,
        generated_at: datetime, signature: str,
    ) -> None:
        sheet = workbook.create_sheet("_核对标识")
        for row, pair in enumerate((
            ("batch_id", batch_id), ("version", version),
            ("generated_at", generated_at.isoformat()), ("signature", signature),
        ), 1):
            sheet.cell(row, 1, pair[0])
            sheet.cell(row, 2, pair[1])
        sheet.sheet_state = "hidden"

    @staticmethod
    def _token(batch_id: str, image_id: str, sequence: int) -> str:
        digest = hashlib.sha256(f"{batch_id}|{image_id}|{sequence}".encode()).hexdigest()[:12]
        return f"{image_id}:{sequence}:{digest}"

    def _signature(self, batch_id: str, version: int, tokens: list[str]) -> str:
        payload = "|".join((batch_id, str(version), *sorted(tokens))).encode()
        return hmac.new(self._secret, payload, hashlib.sha256).hexdigest()

    @staticmethod
    def _resized_image(content: bytes) -> bytes:
        with PillowImage.open(BytesIO(content)) as source:
            source.thumbnail((240, 124), PillowImage.Resampling.LANCZOS)
            output = BytesIO()
            source.convert("RGB").save(output, format="JPEG", quality=85)
            return output.getvalue()

    @staticmethod
    def _issue(code: str, sheet: str, row: int, message: str) -> dict[str, str | int]:
        return {"code": code, "sheet": sheet, "row": row, "message": message}

    @staticmethod
    def _fail(code: str, message: str, sheet: str | None = None) -> None:
        issue: dict[str, str | int] = {"code": code, "message": message}
        if sheet:
            issue["sheet"] = sheet
        raise IncomingWorkbookValidationError([issue])
