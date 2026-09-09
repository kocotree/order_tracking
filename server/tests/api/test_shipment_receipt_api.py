from collections.abc import Iterator

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import Engine
from sqlalchemy.orm import Session, sessionmaker

from app.main import create_app
from app.modules.identity_access import IdentityAccessService
from tests.api.test_shipment_api import ADMIN_ID, ORDER_ID, USER_IDS, _seed


@pytest.fixture
def receipt_clients(
    test_database_engine: Engine, test_database_url: str
) -> Iterator[tuple[TestClient, TestClient, str]]:
    assignment = _seed(test_database_engine, initial_shipped_quantity=5)
    identity = IdentityAccessService(
        sessionmaker(test_database_engine, class_=Session, expire_on_commit=False),
        token_secret=b"receipt-test-token",
        phone_encryption_secret=b"receipt-test-phone",
        phone_digest_secret=b"receipt-test-digest",
    )
    factory_session = identity.issue_session(user_id=USER_IDS[0], terminal="mini")
    admin_session = identity.issue_session(user_id=ADMIN_ID, terminal="web")
    app = create_app(database_url=test_database_url, identity_service=identity)
    with (
        TestClient(app, base_url="https://testserver") as factory,
        TestClient(app, base_url="https://testserver") as admin,
    ):
        factory.headers["Authorization"] = f"Bearer {factory_session.access_token}"
        admin.cookies.set("ot_web_session", admin_session.access_token)
        admin.headers["X-CSRF-Token"] = admin_session.csrf_token
        admin.headers["Idempotency-Key"] = "receipt-confirm-test"
        shipment_id = factory.post(
            "/api/v1/factory/shipments/drafts", json={"preferredOrderId": ORDER_ID}
        ).json()["shipmentId"]
        factory.put(
            f"/api/v1/factory/shipments/drafts/{shipment_id}",
            json={
                "boxes": [
                    {"boxNo": n, "items": [{"assignmentId": assignment, "quantity": q}]}
                    for n, q in [(1, 10), (2, 20)]
                ]
            },
        )
        result = factory.post(
            f"/api/v1/factory/shipments/drafts/{shipment_id}/submit",
            headers={"Idempotency-Key": "receipt-fixture-submit"},
        )
        assert result.status_code == 200, result.text
        yield admin, factory, shipment_id


def test_saved_receipt_reopens_without_changing_published_quantities(
    receipt_clients: tuple[TestClient, TestClient, str],
) -> None:
    admin, factory, shipment_id = receipt_clients
    url = f"/api/v1/admin/shipments/{shipment_id}/receipt"
    response = admin.get(url)
    assert response.status_code == 200, response.text
    draft = response.json()
    assert draft["version"] == 0
    draft["items"][0]["quantity"] = 0
    saved = admin.put(url, json={"version": 0, "items": draft["items"]})
    assert saved.status_code == 200, saved.text
    assert saved.json()["version"] == 1
    assert admin.get(url).json()["items"][0]["quantity"] == 0
    assert admin.get(f"/api/v1/admin/shipments/{shipment_id}").json()["totalQuantity"] == 30
    assert factory.get(f"/api/v1/factory/shipments/{shipment_id}").json()["totalQuantity"] == 30
    assert (
        factory.get("/api/v1/factory/shipment-catalog").json()["items"][0]["shippedQuantity"] == 35
    )
    assert admin.put(url, json={"version": 0, "items": draft["items"]}).status_code == 409


def test_confirm_applies_delta_once_and_preserves_original_export(
    receipt_clients: tuple[TestClient, TestClient, str],
) -> None:
    from io import BytesIO

    from openpyxl import load_workbook

    admin, factory, shipment_id = receipt_clients
    url = f"/api/v1/admin/shipments/{shipment_id}"
    original_export = admin.get(url + "/export").content
    draft = admin.get(url + "/receipt").json()
    draft["items"][0]["quantity"] = 0
    draft["items"][1]["quantity"] = 28
    saved = admin.put(url + "/receipt", json={"version": 0, "items": draft["items"]}).json()
    result = admin.post(url + "/receipt/confirm", json={"version": saved["version"]})
    assert result.status_code == 200, result.text
    assert result.json()["receipt"]["status"] == "CONFIRMED"
    assert result.json()["totalQuantity"] == 28
    assert result.json()["totalBoxes"] == 2
    assert result.json()["receiptDifferences"][0]["quantity"] == -2
    audit = admin.get(f"/api/v1/admin/orders/{ORDER_ID}/audit-logs").json()
    assert any("确认收货，少收 2 件" in item["content"] for item in audit["items"])
    assert admin.post(url + "/receipt/confirm", json={"version": 1}).status_code == 200
    assert (
        factory.get("/api/v1/factory/shipment-catalog").json()["items"][0]["shippedQuantity"] == 33
    )
    assert factory.get(f"/api/v1/factory/shipments/{shipment_id}").json()["totalQuantity"] == 28
    assert (
        admin.get("/api/v1/admin/shipments", params={"orderId": ORDER_ID}).json()["items"][0][
            "totalQuantity"
        ]
        == 28
    )
    assert (
        admin.put(url + "/receipt", json={"version": 1, "items": draft["items"]}).status_code == 409
    )
    before = load_workbook(BytesIO(original_export))
    after = load_workbook(BytesIO(admin.get(url + "/export").content))
    assert [[list(row) for row in sheet.values] for sheet in before] == [
        [list(row) for row in sheet.values] for sheet in after
    ]


@pytest.mark.parametrize("quantity", [0, 28, 35])
def test_confirmed_quantity_cannot_be_autonomously_withdrawn(
    receipt_clients: tuple[TestClient, TestClient, str], quantity: int
) -> None:
    admin, factory, shipment_id = receipt_clients
    url = f"/api/v1/admin/shipments/{shipment_id}"
    draft = admin.get(url + "/receipt").json()
    draft["items"][0]["quantity"] = 0
    draft["items"][1]["quantity"] = quantity
    assert (
        admin.put(url + "/receipt", json={"version": 0, "items": draft["items"]}).status_code == 200
    )
    assert admin.post(url + "/receipt/confirm", json={"version": 1}).status_code == 200
    version = admin.get(url).json()["version"]
    result = factory.post(
        f"/api/v1/factory/shipments/{shipment_id}/withdraw",
        json={"reason": "test", "version": version},
        headers={"Idempotency-Key": "blocked-confirmed"},
    )
    assert result.status_code == 409
    assert (
        factory.get("/api/v1/factory/shipment-catalog").json()["items"][0]["shippedQuantity"]
        == quantity + 5
    )


def test_return_uses_confirmed_base_and_keeps_receipt_difference(
    receipt_clients: tuple[TestClient, TestClient, str],
) -> None:
    admin, factory, shipment_id = receipt_clients
    url = f"/api/v1/admin/shipments/{shipment_id}"
    draft = admin.get(url + "/receipt").json()
    draft["items"][1]["quantity"] = 18
    admin.put(url + "/receipt", json={"version": 0, "items": draft["items"]})
    result = admin.post(url + "/receipt/confirm", json={"version": 1}).json()
    line_id = result["lines"][0]["lineId"]
    assert (
        admin.post(
            url + "/returns",
            json={"reason": "测试退回", "lines": [{"shipmentLineId": line_id, "quantity": 29}]},
        ).status_code
        == 409
    )
    assert (
        admin.post(
            url + "/returns",
            json={"reason": "测试退回", "lines": [{"shipmentLineId": line_id, "quantity": 3}]},
        ).status_code
        == 201
    )
    detail = admin.get(url).json()
    assert detail["lines"][0]["returnableQuantity"] == 25
    assert detail["receiptDifferences"][0]["quantity"] == -2
    assert (
        factory.get("/api/v1/factory/shipment-catalog").json()["items"][0]["shippedQuantity"] == 30
    )


@pytest.mark.parametrize("changed", [False, True])
def test_receipt_notifies_same_factory_once_using_existing_template(
    receipt_clients: tuple[TestClient, TestClient, str], test_database_engine: Engine, changed: bool
) -> None:
    from app.adapters.notifications import FakeWechatNotifier
    from app.modules.notifications_audit import NotificationsAuditService
    from tests.api.test_shipment_api import SAME_FACTORY_USER_ID

    admin, factory, shipment_id = receipt_clients
    service = NotificationsAuditService(
        sessionmaker(test_database_engine, class_=Session, expire_on_commit=False)
    )
    service.record_authorizations(user_id=USER_IDS[0], results={"factory_status": "accepted"})
    url = f"/api/v1/admin/shipments/{shipment_id}"
    draft = admin.get(url + "/receipt").json()
    if changed:
        draft["items"][0]["quantity"] = 12
    admin.put(url + "/receipt", json={"version": 0, "items": draft["items"]})
    while service.consume_next_business_event(worker_id="receipt-test"):
        pass
    assert (
        service.list_notifications(
            user_id=USER_IDS[0], unread_only=False, page=1, page_size=10
        ).total
        == 0
    )
    confirmed = admin.post(url + "/receipt/confirm", json={"version": 1})
    assert confirmed.status_code == 200
    assert admin.post(url + "/receipt/confirm", json={"version": 1}).status_code == 200
    while service.consume_next_business_event(worker_id="receipt-test"):
        pass
    for user in [USER_IDS[0], SAME_FACTORY_USER_ID]:
        page = service.list_notifications(user_id=user, unread_only=False, page=1, page_size=10)
        assert page.total == 1
        assert page.items[0].title == "发货单已确认收货"
        assert page.items[0].target_id == shipment_id
    assert (
        service.list_notifications(
            user_id=USER_IDS[1], unread_only=False, page=1, page_size=10
        ).total
        == 0
    )
    notifier = FakeWechatNotifier()
    while service.deliver_next(worker_id="receipt-test", wechat_notifier=notifier):
        pass
    assert len(notifier.sent) == 1
    assert notifier.sent[0].template_key == "factory_status"
    assert notifier.sent[0].template_data["phrase3"] == "已收货"
    assert notifier.sent[0].target_path.endswith("shipmentId=" + shipment_id)


@pytest.mark.parametrize("invalid", [-1, 1.5, "2", True, 2147483648])
def test_invalid_receipt_quantities_are_rejected(
    receipt_clients: tuple[TestClient, TestClient, str], invalid: object
) -> None:
    admin, _, shipment_id = receipt_clients
    url = f"/api/v1/admin/shipments/{shipment_id}/receipt"
    draft = admin.get(url).json()
    draft["items"][0]["quantity"] = invalid
    assert admin.put(url, json={"version": 0, "items": draft["items"]}).status_code == 422
    assert admin.get(url).json()["version"] == 0


@pytest.mark.parametrize("blocked", ["return", "pending", "voided"])
def test_inverse_state_blocks_receipt_writes(
    receipt_clients: tuple[TestClient, TestClient, str], blocked: str,
    test_database_engine: Engine,
) -> None:
    admin, factory, shipment_id = receipt_clients
    url = f"/api/v1/admin/shipments/{shipment_id}"
    draft = admin.get(url + "/receipt").json()
    if blocked == "return":
        line_id = admin.get(url).json()["lines"][0]["lineId"]
        assert (
            admin.post(
                url + "/returns",
                json={"reason": "测试", "lines": [{"shipmentLineId": line_id, "quantity": 1}]},
            ).status_code
            == 201
        )
    else:
        from sqlalchemy.orm import Session

        from app.db.models import Shipment
        with Session(test_database_engine) as session, session.begin():
            stored = session.get(Shipment, shipment_id)
            assert stored is not None
            stored.status = "VOIDED" if blocked == "voided" else "VOID_PENDING"
    assert (
        admin.put(url + "/receipt", json={"version": 0, "items": draft["items"]}).status_code == 409
    )
    assert admin.post(url + "/receipt/confirm", json={"version": 0}).status_code == 409


def test_original_can_be_confirmed_without_draft_and_stale_version_cannot_confirm(
    receipt_clients: tuple[TestClient, TestClient, str],
) -> None:
    admin, factory, shipment_id = receipt_clients
    url = f"/api/v1/admin/shipments/{shipment_id}"
    result = admin.post(url + "/receipt/confirm", json={"version": 0})
    assert result.status_code == 200, result.text
    assert result.json()["receiptDifferences"] == []
    assert admin.post(url + "/receipt/confirm", json={"version": 0}).status_code == 200
    assert (
        factory.get("/api/v1/factory/shipment-catalog").json()["items"][0]["shippedQuantity"] == 35
    )


def test_receipt_requires_web_admin_and_csrf(
    receipt_clients: tuple[TestClient, TestClient, str],
) -> None:
    admin, factory, shipment_id = receipt_clients
    url = f"/api/v1/admin/shipments/{shipment_id}/receipt"
    draft = admin.get(url).json()
    assert factory.get(url).status_code == 401
    assert factory.put(url, json={"version": 0, "items": draft["items"]}).status_code == 401
    assert factory.post(url + "/confirm", json={"version": 0}).status_code == 401
    del admin.headers["X-CSRF-Token"]
    assert admin.put(url, json={"version": 0, "items": draft["items"]}).status_code == 403
    assert admin.post(url + "/confirm", json={"version": 0}).status_code == 403


def test_receipt_rejects_changed_item_set_and_stale_confirmation(
    receipt_clients: tuple[TestClient, TestClient, str],
) -> None:
    admin, _, shipment_id = receipt_clients
    url = f"/api/v1/admin/shipments/{shipment_id}/receipt"
    items = admin.get(url).json()["items"]
    assert admin.put(url, json={"version": 0, "items": [items[0], items[0]]}).status_code == 422
    assert admin.put(url, json={"version": 0, "items": items[:1]}).status_code == 422
    assert admin.put(url, json={"version": 0, "items": items}).status_code == 200
    assert admin.post(url + "/confirm", json={"version": 0}).status_code == 409


def test_short_receipt_reopens_completed_order_and_preserves_initial_baseline(
    receipt_clients: tuple[TestClient, TestClient, str],
) -> None:
    admin, factory, shipment_id = receipt_clients
    url = f"/api/v1/admin/shipments/{shipment_id}"
    assert admin.post(f"/api/v1/admin/orders/{ORDER_ID}/complete").status_code == 200
    draft = admin.get(url + "/receipt").json()
    draft["items"][1]["quantity"] = 18
    assert (
        admin.put(url + "/receipt", json={"version": 0, "items": draft["items"]}).status_code == 200
    )
    assert admin.post(url + "/receipt/confirm", json={"version": 1}).status_code == 200
    order = admin.get(f"/api/v1/orders/{ORDER_ID}").json()
    assert order["lifecycle"] == "PUBLISHED"
    assert (
        factory.get("/api/v1/factory/shipment-catalog").json()["items"][0]["shippedQuantity"] == 33
    )


def test_concurrent_confirmations_count_once(
    receipt_clients: tuple[TestClient, TestClient, str],
) -> None:
    from concurrent.futures import ThreadPoolExecutor

    admin, factory, shipment_id = receipt_clients
    url = f"/api/v1/admin/shipments/{shipment_id}/receipt"
    draft = admin.get(url).json()
    draft["items"][0]["quantity"] = 0
    admin.put(url, json={"version": 0, "items": draft["items"]})
    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(
            pool.map(lambda _: admin.post(url + "/confirm", json={"version": 1}), range(2))
        )
    assert [result.status_code for result in results] == [200, 200]
    assert (
        factory.get("/api/v1/factory/shipment-catalog").json()["items"][0]["shippedQuantity"] == 25
    )


def test_concurrent_draft_saves_conflict_instead_of_overwriting(
    receipt_clients: tuple[TestClient, TestClient, str],
) -> None:
    from concurrent.futures import ThreadPoolExecutor

    admin, _, shipment_id = receipt_clients
    url = f"/api/v1/admin/shipments/{shipment_id}/receipt"
    items = admin.get(url).json()["items"]
    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(
            pool.map(lambda _: admin.put(url, json={"version": 0, "items": items}), range(2))
        )
    assert sorted(result.status_code for result in results) == [200, 409]
    assert admin.get(url).json()["version"] == 1


def test_confirmation_database_failure_rolls_back_every_business_effect(
    receipt_clients: tuple[TestClient, TestClient, str], test_database_engine: Engine
) -> None:
    from sqlalchemy import event

    admin, factory, shipment_id = receipt_clients
    url = f"/api/v1/admin/shipments/{shipment_id}"
    draft = admin.get(url + "/receipt").json()
    draft["items"][0]["quantity"] = 0
    admin.put(url + "/receipt", json={"version": 0, "items": draft["items"]})

    def fail_outbox(connection, cursor, statement, parameters, context, executemany):
        if statement.startswith("INSERT INTO outbox_messages"):
            raise RuntimeError("simulated transaction failure")

    # Engine class listener reaches the API's independent pool, at the database boundary.
    event.listen(Engine, "before_cursor_execute", fail_outbox)
    try:
        with pytest.raises(RuntimeError, match="simulated transaction failure"):
            admin.post(url + "/receipt/confirm", json={"version": 1})
    finally:
        event.remove(Engine, "before_cursor_execute", fail_outbox)
    assert admin.get(url + "/receipt").json()["status"] == "DRAFT"
    assert admin.get(url).json()["totalQuantity"] == 30
    assert (
        factory.get("/api/v1/factory/shipment-catalog").json()["items"][0]["shippedQuantity"] == 35
    )
    assert admin.post(url + "/receipt/confirm", json={"version": 1}).status_code == 200
    assert (
        factory.get("/api/v1/factory/shipment-catalog").json()["items"][0]["shippedQuantity"] == 25
    )


def test_receipt_and_return_race_serializes_quantity_and_guard(
    receipt_clients: tuple[TestClient, TestClient, str],
) -> None:
    from concurrent.futures import ThreadPoolExecutor

    admin, factory, shipment_id = receipt_clients
    url = f"/api/v1/admin/shipments/{shipment_id}"
    draft = admin.get(url + "/receipt").json()
    draft["items"][0]["quantity"] = 0
    admin.put(url + "/receipt", json={"version": 0, "items": draft["items"]})
    line_id = admin.get(url).json()["lines"][0]["lineId"]
    with ThreadPoolExecutor(max_workers=2) as pool:
        confirm = pool.submit(admin.post, url + "/receipt/confirm", json={"version": 1})
        returned = pool.submit(
            admin.post,
            url + "/returns",
            json={"reason": "测试竞争", "lines": [{"shipmentLineId": line_id, "quantity": 3}]},
        )
        confirmation_status = confirm.result().status_code
        assert returned.result().status_code == 201
    assert confirmation_status in (200, 409)
    expected = 22 if confirmation_status == 200 else 32
    assert (
        factory.get("/api/v1/factory/shipment-catalog").json()["items"][0]["shippedQuantity"]
        == expected
    )


def test_mixed_box_same_sku_keeps_each_orders_difference(
    receipt_clients: tuple[TestClient, TestClient, str], test_database_engine: Engine
) -> None:
    from datetime import date

    from app.modules.orders import AssignmentInput, DraftLineInput, OrderService
    from tests.api.test_shipment_api import FACTORY_IDS, VARIANT_ID

    admin, factory, _ = receipt_clients
    service = OrderService(
        sessionmaker(test_database_engine, class_=Session, expire_on_commit=False)
    )
    second = service.create_draft(
        actor_id=ADMIN_ID,
        order_no="RECEIPT-B",
        order_date=date(2026, 9, 7),
        tracker="松子",
        lines=[
            DraftLineInput(
                variant_id=VARIANT_ID,
                order_quantity=50,
                assignments=[
                    AssignmentInput(
                        factory_id=FACTORY_IDS[0], quantity=50, contract_ship_date=date(2026, 9, 10)
                    )
                ],
            )
        ],
        request_id="receipt-second",
    )
    service.publish(
        actor_id=ADMIN_ID,
        order_id=second.order_id,
        version=second.version,
        request_id="receipt-second-publish",
        idempotency_key="receipt-second-publish",
    )
    catalog = factory.get("/api/v1/factory/shipment-catalog").json()["items"]
    assignments = {item["orderId"]: item["assignmentId"] for item in catalog}
    shipment_id = factory.post("/api/v1/factory/shipments/drafts", json={}).json()["shipmentId"]
    assert (
        factory.put(
            f"/api/v1/factory/shipments/drafts/{shipment_id}",
            json={
                "boxes": [
                    {
                        "boxNo": 1,
                        "items": [
                            {"assignmentId": assignments[ORDER_ID], "quantity": 10},
                            {"assignmentId": assignments[second.order_id], "quantity": 20},
                        ],
                    }
                ]
            },
        ).status_code
        == 200
    )
    assert (
        factory.post(
            f"/api/v1/factory/shipments/drafts/{shipment_id}/submit",
            headers={"Idempotency-Key": "mixed-submit"},
        ).status_code
        == 200
    )
    url = f"/api/v1/admin/shipments/{shipment_id}/receipt"
    draft = admin.get(url).json()
    draft["items"][0]["quantity"] = 0
    draft["items"][1]["quantity"] = 25
    admin.put(url, json={"version": 0, "items": draft["items"]})
    result = admin.post(url + "/confirm", json={"version": 1})
    assert result.status_code == 200, result.text
    assert result.json()["totalQuantity"] == 25
    assert {item["orderNo"]: item["quantity"] for item in result.json()["receiptDifferences"]} == {
        "S07-ORDER-A": -10,
        "RECEIPT-B": 5,
    }
    catalog = factory.get("/api/v1/factory/shipment-catalog").json()["items"]
    assert {item["orderId"]: item["shippedQuantity"] for item in catalog} == {
        ORDER_ID: 35,
        second.order_id: 25,
    }


def test_receipt_delivery_rechecks_factory_membership(
    receipt_clients: tuple[TestClient, TestClient, str], test_database_engine: Engine
) -> None:
    from app.adapters.notifications import FakeWechatNotifier
    from app.db.models import User
    from app.modules.notifications_audit import NotificationsAuditService
    from tests.api.test_shipment_api import FACTORY_IDS

    admin, _, shipment_id = receipt_clients
    service = NotificationsAuditService(
        sessionmaker(test_database_engine, class_=Session, expire_on_commit=False)
    )
    service.record_authorizations(user_id=USER_IDS[0], results={"factory_status": "accepted"})
    assert (
        admin.post(
            f"/api/v1/admin/shipments/{shipment_id}/receipt/confirm", json={"version": 0}
        ).status_code
        == 200
    )
    while service.consume_next_business_event(worker_id="receipt-test"):
        pass
    with Session(test_database_engine) as session, session.begin():
        session.get(User, USER_IDS[0]).factory_id = FACTORY_IDS[1]
    notifier = FakeWechatNotifier()
    while service.deliver_next(worker_id="receipt-test", wechat_notifier=notifier):
        pass
    assert notifier.sent == []
