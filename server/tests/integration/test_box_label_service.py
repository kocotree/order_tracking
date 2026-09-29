from concurrent.futures import ThreadPoolExecutor
from datetime import UTC, datetime
from io import BytesIO

import pytest
from openpyxl import load_workbook
from PIL import Image
from sqlalchemy import Engine
from sqlalchemy.orm import Session, sessionmaker

from app.adapters.private_files import (
    FakePrivateFileStore,
    PrivateFileNotFound,
    PrivateFileStoreUnavailable,
)
from app.db.models import (
    BoxLabelExport,
    Factory,
    Order,
    OrderAssignment,
    OrderDetail,
    OrderLine,
    Product,
    ProductVariant,
    User,
)
from app.modules.box_labels import BoxLabelService
from app.modules.box_labels.service import BoxLabelGenerationError
from app.modules.box_labels.workbook import BoxLabelWorkbookRenderer
from tests.integration.test_contracts import (
    ADMIN_ID,
    FACTORY_ID,
    ORDER_ID,
    PRODUCT_ID,
    VARIANT_ID,
    _seed_published_order,
)


def _picture() -> bytes:
    output = BytesIO()
    Image.new("RGB", (20, 10), "blue").save(output, "PNG")
    return output.getvalue()


def test_group_export_is_stable_and_fails_before_freezing_unavailable_image(
    test_database_engine: Engine,
) -> None:
    _seed_published_order(test_database_engine)
    sessions = sessionmaker(test_database_engine, class_=Session, expire_on_commit=False)
    files = FakePrivateFileStore(bucket="box-label-test")
    service = BoxLabelService(
        sessions, renderer=BoxLabelWorkbookRenderer(), file_store=files
    )
    with sessions() as session, session.begin():
        original = session.get(ProductVariant, VARIANT_ID)
        assert original is not None
        original.properties_value = "米色 / 52cm"
        original.image_source_ref = "source-not-cached"
        original.image_cache_status = "pending"
        second = ProductVariant(
            variant_id="box-label-second-variant", product_id=PRODUCT_ID,
            source_sku_id="6970000000002", properties_value="米色 / 54cm",
            source_category="童帽春夏", source_enabled=1, is_available=True,
            image_source_ref="second-image", image_object_key="images/second.png",
            image_cache_status="cached", source_modified_at=original.source_modified_at,
            first_synced_at=original.first_synced_at,
            last_synced_at=original.last_synced_at,
        )
        session.add(second)
        session.flush()
        order = session.get(Order, ORDER_ID)
        assert order is not None
        line = OrderLine(
            order_id=ORDER_ID, product_variant_id=second.variant_id,
            order_quantity=30, sku_id_snapshot=second.source_sku_id,
            product_name_snapshot="儿童遮阳帽", properties_value_snapshot="米色 / 54cm",
            category_snapshot="童帽春夏",
        )
        session.add(line)
        session.flush()
        session.add(OrderAssignment(
            order_line_id=line.order_line_id, factory_id=FACTORY_ID,
            assigned_quantity=30, factory_name_snapshot="合同测试工厂",
        ))
        session.add(User(
            user_id="box-label-second-admin", role="admin", is_enabled=True,
            feishu_display_name="测试管理员",
        ))
    groups = service.list_for_order(actor_id=ADMIN_ID, order_id=ORDER_ID)
    assert len(groups) == 1
    assert groups[0].color == "米色"
    assert groups[0].eligible
    with pytest.raises(PrivateFileNotFound):
        service.export(
            actor_id=ADMIN_ID, order_id=ORDER_ID, selected_group_id=groups[0].group_id
        )
    files.put(object_key="images/second.png", content=_picture(), content_type="image/png")
    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(pool.map(
            lambda actor: service.export(
                actor_id=actor, order_id=ORDER_ID, selected_group_id=groups[0].group_id
            ),
            [ADMIN_ID, "box-label-second-admin"],
        ))
    assert results[0] == results[1]
    export_id, filename = results[0]
    assert filename == "HT-ORDER-001_合同测试工厂_儿童遮阳帽_米色.xlsx"
    assert files.object_count == 2
    downloaded = service.download(actor_id=ADMIN_ID, export_id=export_id)[1]
    sheet = load_workbook(BytesIO(downloaded)).active
    assert sheet is not None
    assert len(sheet._images) == 0
    assert sheet["A2"].value == sheet["A12"].value == "#VALUE!"
    assert sheet["B1"].value == sheet["B11"].value == "工厂：合同测试工厂"
    assert sheet["B4"].value == sheet["B14"].value == "颜色：米色"
    assert sheet["B5"].value == sheet["B15"].value == "尺码："
    with sessions() as session:
        saved = session.get(BoxLabelExport, export_id)
        assert saved is not None and saved.template_version == "v2"
    with sessions() as session, session.begin():
        factory = session.get(Factory, FACTORY_ID)
        product = session.get(Product, PRODUCT_ID)
        assert factory is not None and product is not None
        factory.factory_name = "改名工厂"
        product.name = "改名产品"
    assert service.export(
        actor_id="box-label-second-admin", order_id=ORDER_ID,
        selected_group_id=groups[0].group_id,
    ) == results[0]
    assert service.download(actor_id="box-label-second-admin", export_id=export_id)[1] == downloaded


def test_true_missing_image_is_blank_but_pending_image_is_blocked(
    test_database_engine: Engine,
) -> None:
    _seed_published_order(test_database_engine)
    sessions = sessionmaker(test_database_engine, class_=Session, expire_on_commit=False)
    files = FakePrivateFileStore(bucket="box-label-blank-test")
    service = BoxLabelService(
        sessions,
        renderer=BoxLabelWorkbookRenderer(),
        file_store=files,
    )
    group = service.list_for_order(actor_id=ADMIN_ID, order_id=ORDER_ID)[0]
    with sessions() as session, session.begin():
        variant = session.get(ProductVariant, VARIANT_ID)
        assert variant is not None
        variant.image_source_ref = "uncached-source"
        variant.image_cache_status = "pending"
    with pytest.raises(BoxLabelGenerationError):
        service.export(
            actor_id=ADMIN_ID, order_id=ORDER_ID, selected_group_id=group.group_id
        )
    assert files.object_count == 0
    with sessions() as session, session.begin():
        variant = session.get(ProductVariant, VARIANT_ID)
        assert variant is not None
        variant.image_source_ref = None
        variant.image_cache_status = "missing"
    export_id, _ = service.export(
        actor_id=ADMIN_ID, order_id=ORDER_ID, selected_group_id=group.group_id
    )
    content = service.download(actor_id=ADMIN_ID, export_id=export_id)[1]
    sheet = load_workbook(BytesIO(content)).active
    assert sheet is not None
    assert len(sheet._images) == 0


def test_matched_factory_with_unmatched_product_or_color_remains_visible(
    test_database_engine: Engine,
) -> None:
    _seed_published_order(test_database_engine)
    sessions = sessionmaker(test_database_engine, class_=Session, expire_on_commit=False)
    with sessions() as session, session.begin():
        order = session.get(Order, ORDER_ID)
        assert order is not None
        order.detail_mode = True
        now = datetime.now(UTC)
        session.add_all([
            OrderDetail(
                detail_id="box-label-unmatched-product", order_id=ORDER_ID,
                origin="manual", sort_order=1, accepted_raw_fields={},
                product_name="未匹配产品", properties_value="冰川蓝110",
                matched_factory_id=FACTORY_ID, dispatch_state="UNASSIGNED",
                parse_issues=[], created_at=now, updated_at=now,
            ),
            OrderDetail(
                detail_id="box-label-unmatched-color", order_id=ORDER_ID,
                origin="manual", sort_order=2, accepted_raw_fields={},
                product_name="儿童遮阳帽", properties_value="蓝2",
                matched_factory_id=FACTORY_ID, matched_variant_id=VARIANT_ID,
                dispatch_state="UNASSIGNED", parse_issues=[],
                created_at=now, updated_at=now,
            ),
        ])
    service = BoxLabelService(
        sessions,
        renderer=BoxLabelWorkbookRenderer(),
        file_store=FakePrivateFileStore(bucket="box-label-unmatched-test"),
    )
    rows = service.list_for_order(actor_id=ADMIN_ID, order_id=ORDER_ID)
    assert [row.ineligible_reason for row in rows] == ["产品未匹配", "颜色无法确定"]
    assert all(not row.eligible for row in rows)


def test_failed_storage_does_not_freeze_the_first_export(test_database_engine: Engine) -> None:
    _seed_published_order(test_database_engine)
    sessions = sessionmaker(test_database_engine, class_=Session, expire_on_commit=False)
    renderer = BoxLabelWorkbookRenderer()
    broken = BoxLabelService(
        sessions, renderer=renderer,
        file_store=FakePrivateFileStore(bucket="box-label-retry", fail_put=True),
    )
    group = broken.list_for_order(actor_id=ADMIN_ID, order_id=ORDER_ID)[0]
    with pytest.raises(PrivateFileStoreUnavailable):
        broken.export(actor_id=ADMIN_ID, order_id=ORDER_ID, selected_group_id=group.group_id)
    healthy = BoxLabelService(
        sessions, renderer=renderer,
        file_store=FakePrivateFileStore(bucket="box-label-retry"),
    )
    export_id, _ = healthy.export(
        actor_id=ADMIN_ID, order_id=ORDER_ID, selected_group_id=group.group_id
    )
    assert healthy.download(actor_id=ADMIN_ID, export_id=export_id)[1].startswith(b"PK")
