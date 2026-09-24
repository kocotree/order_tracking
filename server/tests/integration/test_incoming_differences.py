"""Slice 1 acceptance for incoming differences: AC-08/09/11/12/13/17 plus batch_no and matching."""

import hashlib
import json
from collections.abc import Iterator
from datetime import UTC, date, datetime, timedelta
from io import BytesIO
from typing import Any

import pytest
from openpyxl import load_workbook
from PIL import Image
from sqlalchemy import Engine, delete, func, select
from sqlalchemy.orm import Session, sessionmaker

from app.adapters.notifications import (
    DeliveryRequest,
    FakeWechatNotifier,
    NotificationDeliveryError,
)
from app.adapters.private_files import FakePrivateFileStore
from app.adapters.vision import FakeIncomingDiffRecognizer
from app.db.models import (
    AuditLog,
    ExternalIdentity,
    Factory,
    IdempotencyRecord,
    IncomingDiffAdjustment,
    IncomingDiffBatch,
    IncomingDiffImage,
    IncomingDiffRecord,
    IncomingDiffWorkbook,
    Notification,
    NotificationAuthorization,
    Order,
    OrderAssignment,
    OrderDetail,
    OrderLine,
    OutboxMessage,
    Product,
    ProductVariant,
    QuantityLedger,
    StoredFile,
    User,
)
from app.modules.incoming_differences import (
    IncomingDifferenceConflict,
    IncomingDifferencePermissionDenied,
    IncomingDifferenceService,
    IncomingDifferenceValidationError,
)
from app.modules.incoming_differences.bot import CONFIRM_JOB, FeishuBotService
from app.modules.incoming_differences.recognition import IncomingDiffRecognitionService
from app.modules.incoming_differences.workbook import IncomingWorkbookCodec
from app.modules.infrastructure import InfrastructureStore, utc_now
from app.modules.notifications_audit import NotificationsAuditService
from app.settings.config import Settings
from app.worker.runtime import Worker

ADMIN = "idf-admin"
OTHER_ADMIN = "idf-admin-2"
DISABLED_ADMIN = "idf-admin-off"
FACTORY_USER = "idf-factory-user"
SECOND_FACTORY_USER = "idf-factory-user-2"
DISABLED_FACTORY_USER = "idf-factory-user-off"
USER_IDS = [
    ADMIN,
    OTHER_ADMIN,
    DISABLED_ADMIN,
    FACTORY_USER,
    SECOND_FACTORY_USER,
    DISABLED_FACTORY_USER,
]
FACTORY_A = "idf-factory-a"
FACTORY_B = "idf-factory-b"
FACTORY_IDS = [FACTORY_A, FACTORY_B]
ORDER_ID = "idf-order"
BATCH_ID = "idf-batch"
WORKBOOK_ID = "idf-workbook"
IMAGE_ID = "idf-image"
FILE_IDS = [9601, 9602]
NOW = datetime(2026, 9, 18, 2, 30, tzinfo=UTC)
SOURCE_TIME = datetime(2026, 9, 1, 8, 0)


def _clean(engine: Engine) -> None:
    with Session(engine) as session, session.begin():
        session.execute(
            delete(IncomingDiffAdjustment).where(
                IncomingDiffAdjustment.actor_id.in_(USER_IDS)
            )
        )
        session.execute(
            delete(IncomingDiffRecord).where(IncomingDiffRecord.order_id == ORDER_ID)
        )
        session.execute(
            delete(QuantityLedger).where(QuantityLedger.actor_id.in_(USER_IDS))
        )
        session.execute(
            delete(Notification).where(Notification.recipient_id.in_(USER_IDS))
        )
        session.execute(delete(NotificationAuthorization).where(
            NotificationAuthorization.user_id.in_(USER_IDS)
        ))
        session.execute(
            delete(OutboxMessage).where(
                OutboxMessage.dedupe_key.like("%incoming-diff:%")
            )
        )
        session.execute(delete(OutboxMessage).where(
            OutboxMessage.event_type == "incoming_diff.bot_reply"
        ))
        session.execute(
            delete(AuditLog).where(AuditLog.action.like("incoming_diff_%"))
        )
        session.execute(
            delete(IdempotencyRecord).where(
                IdempotencyRecord.scope == "incoming_diff_confirm"
            )
        )
        session.execute(delete(IdempotencyRecord).where(
            IdempotencyRecord.scope == "feishu_event"
        ))
        session.execute(
            IncomingDiffBatch.__table__.update().values(current_workbook_id=None)
        )
        session.execute(delete(IncomingDiffImage))
        session.execute(delete(IncomingDiffWorkbook))
        session.execute(delete(IncomingDiffBatch))
        session.execute(
            OrderAssignment.__table__.update()
            .where(
                OrderAssignment.detail_id.in_(
                    select(OrderDetail.detail_id).where(OrderDetail.order_id == ORDER_ID)
                )
            )
            .values(detail_id=None)
        )
        session.execute(delete(OrderDetail).where(OrderDetail.order_id == ORDER_ID))
        session.execute(
            delete(OrderAssignment).where(
                OrderAssignment.order_line_id.in_(
                    select(OrderLine.order_line_id).where(OrderLine.order_id == ORDER_ID)
                )
            )
        )
        session.execute(delete(OrderLine).where(OrderLine.order_id == ORDER_ID))
        session.execute(delete(Order).where(Order.order_id == ORDER_ID))
        session.execute(OrderAssignment.__table__.update().where(
            OrderAssignment.detail_id == "idf-detail-later"
        ).values(detail_id=None))
        session.execute(delete(OrderDetail).where(OrderDetail.order_id == "idf-order-2"))
        session.execute(delete(OrderAssignment).where(
            OrderAssignment.order_line_id.in_(select(OrderLine.order_line_id).where(
                OrderLine.order_id == "idf-order-2"
            ))
        ))
        session.execute(delete(OrderLine).where(OrderLine.order_id == "idf-order-2"))
        session.execute(delete(Order).where(Order.order_id == "idf-order-2"))
        session.execute(delete(StoredFile).where(StoredFile.file_id.in_((9709, 9710))))
        session.execute(delete(StoredFile).where(StoredFile.file_id.in_(FILE_IDS)))
        session.execute(delete(StoredFile).where(
            StoredFile.object_key.like("incoming-differences/%")
        ))
        session.execute(delete(ExternalIdentity).where(
            ExternalIdentity.user_id.in_(USER_IDS)
        ))
        session.execute(
            delete(ProductVariant).where(
                ProductVariant.variant_id.in_(["idf-variant-1", "idf-variant-2"])
            )
        )
        session.execute(delete(Product).where(Product.product_id == "idf-product"))
        session.execute(delete(User).where(User.user_id.in_(USER_IDS)))
        session.execute(delete(Factory).where(Factory.factory_id.in_(FACTORY_IDS)))


@pytest.fixture(autouse=True)
def clean_incoming_differences(test_database_engine: Engine) -> Iterator[None]:
    _clean(test_database_engine)
    yield
    _clean(test_database_engine)


def _seed_masters(session: Session) -> None:
    session.add_all(
        [
            Factory(
                factory_id=factory_id,
                supplier_number=f"IDF-{index}",
                factory_name=f"来货测试工厂{index}",
                factory_code=f"IDF{index}",
                is_enabled=True,
            )
            for index, factory_id in enumerate(FACTORY_IDS, 1)
        ]
    )
    session.flush()
    session.add_all(
        [
            User(user_id=ADMIN, role="admin", is_enabled=True, feishu_display_name="来货管理员"),
            User(
                user_id=OTHER_ADMIN,
                role="admin",
                is_enabled=True,
                feishu_display_name="另一管理员",
            ),
            User(
                user_id=DISABLED_ADMIN,
                role="admin",
                is_enabled=False,
                feishu_display_name="停用管理员",
            ),
            User(
                user_id=FACTORY_USER,
                role="factory",
                is_enabled=True,
                feishu_display_name="工厂用户",
                factory_id=FACTORY_A,
                factory_position="employee",
            ),
        ]
    )
    session.add(
        Product(
            product_id="idf-product",
            source_i_id="IDF-ITEM",
            name="来货测试产品",
            is_available=True,
            image_cache_status="missing",
            source_modified_at=SOURCE_TIME,
            first_synced_at=SOURCE_TIME,
            last_synced_at=SOURCE_TIME,
        )
    )
    session.flush()
    session.add_all(
        [
            ProductVariant(
                variant_id="idf-variant-1",
                product_id="idf-product",
                source_sku_id="IDF-SKU-1",
                properties_value="红色 / 110",
                is_available=True,
                source_modified_at=SOURCE_TIME,
                first_synced_at=SOURCE_TIME,
                last_synced_at=SOURCE_TIME,
            ),
            ProductVariant(
                variant_id="idf-variant-2",
                product_id="idf-product",
                source_sku_id="IDF-SKU-2",
                properties_value="蓝色 / 120",
                is_available=True,
                source_modified_at=SOURCE_TIME,
                first_synced_at=SOURCE_TIME,
                last_synced_at=SOURCE_TIME,
            ),
        ]
    )
    session.add_all(
        [
            StoredFile(
                file_id=FILE_IDS[0],
                bucket="incoming-diff",
                object_key="images/idf/1.jpg",
                original_filename="来货.jpg",
                mime_type="image/jpeg",
                size_bytes=1024,
                content_sha256="a" * 64,
                uploaded_by=ADMIN,
            ),
            StoredFile(
                file_id=FILE_IDS[1],
                bucket="incoming-diff",
                object_key="workbooks/idf/1.xlsx",
                original_filename="核对表.xlsx",
                mime_type=(
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                ),
                size_bytes=2048,
                content_sha256="b" * 64,
                uploaded_by=ADMIN,
            ),
        ]
    )


def _seed_order(session: Session) -> tuple[int, int]:
    session.add(
        Order(
            order_id=ORDER_ID,
            order_no="IDF-ORDER-1",
            source="manual",
            order_date=date(2026, 9, 1),
            tracker="松子",
            trackers=["松子"],
            detail_mode=True,
            lifecycle="PUBLISHED",
            created_by=ADMIN,
            updated_by=ADMIN,
            created_at=SOURCE_TIME,
            updated_at=SOURCE_TIME,
        )
    )
    session.flush()
    lines = []
    for index, (variant_id, properties, quantity) in enumerate(
        [("idf-variant-1", "红色 / 110", 100), ("idf-variant-2", "蓝色 / 120", 50)], 1
    ):
        line = OrderLine(
            order_id=ORDER_ID,
            product_variant_id=variant_id,
            order_quantity=quantity,
            sku_id_snapshot=f"IDF-SKU-{index}",
            product_name_snapshot="来货测试产品",
            properties_value_snapshot=properties,
            created_at=SOURCE_TIME,
            updated_at=SOURCE_TIME,
        )
        session.add(line)
        lines.append(line)
    session.flush()
    assignments = []
    for index, (line, factory_id, initial_shipped) in enumerate(
        [(lines[0], FACTORY_A, 10), (lines[1], FACTORY_B, 0)], 1
    ):
        assignment = OrderAssignment(
            order_line_id=line.order_line_id,
            factory_id=factory_id,
            is_active=True,
            contract_ship_date=date(2026, 9, 20),
            assigned_quantity=line.order_quantity,
            initial_shipped_quantity=initial_shipped,
            factory_name_snapshot=f"来货测试工厂{index}",
            created_at=SOURCE_TIME,
            updated_at=SOURCE_TIME,
        )
        session.add(assignment)
        assignments.append(assignment)
    session.flush()
    for index, (line, assignment) in enumerate(zip(lines, assignments, strict=True), 1):
        detail_id = f"idf-detail-{index}"
        session.add(
            OrderDetail(
                detail_id=detail_id,
                order_id=ORDER_ID,
                origin="manual",
                sort_order=index,
                accepted_raw_fields={},
                purchase_order_id=f"PO-{index}",
                purchase_order_item_id=f"POI-{index}",
                source_sku_id=line.sku_id_snapshot,
                product_name="来货测试产品",
                properties_value=line.properties_value_snapshot,
                matched_variant_id=line.product_variant_id,
                matched_factory_id=assignment.factory_id,
                order_quantity=line.order_quantity,
                source_trackers=["松子"],
                contract_ship_date=date(2026, 9, 20),
                parse_issues=[],
                assignment_id=assignment.order_assignment_id,
                dispatch_state="ASSIGNED",
                created_at=SOURCE_TIME,
                updated_at=SOURCE_TIME,
            )
        )
    session.flush()
    for index, assignment in enumerate(assignments, 1):
        assignment.detail_id = f"idf-detail-{index}"
    session.flush()
    return assignments[0].order_assignment_id, assignments[1].order_assignment_id


def _seed_batch(
    session: Session,
    *,
    lines: list[dict[str, Any]],
    status: str = "READY",
    submitter_id: str = ADMIN,
    version: int = 1,
) -> None:
    session.add(
        IncomingDiffBatch(
            batch_id=BATCH_ID,
            batch_no="IN20260918-01",
            submitter_id=submitter_id,
            feishu_chat_id="chat-1",
            feishu_open_id="open-1",
            status=status,
            created_at=SOURCE_TIME,
            updated_at=SOURCE_TIME,
            frozen_at=SOURCE_TIME,
        )
    )
    session.flush()
    session.add(
        IncomingDiffImage(
            image_id=IMAGE_ID,
            batch_id=BATCH_ID,
            sort_order=1,
            feishu_message_id="msg-1",
            feishu_image_key="img-key-1",
            file_id=FILE_IDS[0],
            content_sha256="a" * 64,
            ocr_status="SUCCEEDED",
            created_at=SOURCE_TIME,
        )
    )
    session.add(
        IncomingDiffWorkbook(
            workbook_id=WORKBOOK_ID,
            batch_id=BATCH_ID,
            version=version,
            direction="UPLOADED",
            file_id=FILE_IDS[1],
            content_sha256="b" * 64,
            line_snapshot=lines,
            submitted_by=submitter_id,
            submitted_at=SOURCE_TIME,
        )
    )
    session.flush()
    batch = session.get(IncomingDiffBatch, BATCH_ID)
    assert batch is not None
    batch.current_workbook_id = WORKBOOK_ID


def _line(
    assignment_id: int,
    quantity: int,
    *,
    purchase_order_id: str | None = "PO-1",
    purchase_order_item_id: str | None = "POI-1",
    row_number: int = 2,
    sheet_name: str = "来货测试工厂1",
    image_id: str = IMAGE_ID,
) -> dict[str, Any]:
    return {
        "imageId": image_id,
        "sheetName": sheet_name,
        "rowNumber": row_number,
        "orderAssignmentId": assignment_id,
        "purchaseOrderId": purchase_order_id,
        "purchaseOrderItemId": purchase_order_item_id,
        "quantity": quantity,
    }


def _service(engine: Engine) -> IncomingDifferenceService:
    sessions = sessionmaker(engine, class_=Session, expire_on_commit=False)
    return IncomingDifferenceService(sessions, clock=lambda: NOW)


def _unified_shipped(session: Session, assignment_id: int) -> int:
    assignment = session.get(OrderAssignment, assignment_id)
    assert assignment is not None
    ledger_total = int(
        session.scalar(
            select(func.coalesce(func.sum(QuantityLedger.quantity_delta), 0)).where(
                QuantityLedger.order_assignment_id == assignment_id
            )
        )
        or 0
    )
    return assignment.initial_shipped_quantity + ledger_total


def _confirmed_record(engine: Engine, *, quantity: int = 5) -> tuple[int, str]:
    with Session(engine) as session, session.begin():
        _seed_masters(session)
        first, _second = _seed_order(session)
        _seed_batch(session, lines=[_line(first, quantity)])
    _service(engine).confirm(batch_id=BATCH_ID, workbook_version=1, actor_id=ADMIN)
    with Session(engine) as session:
        record_id = session.scalar(
            select(IncomingDiffRecord.record_id).where(
                IncomingDiffRecord.order_assignment_id == first
            )
        )
        assert record_id is not None
    return first, record_id


def test_confirm_writes_records_ledger_audit_and_one_outbox_per_factory_and_order(
    test_database_engine: Engine,
) -> None:
    with Session(test_database_engine) as session, session.begin():
        _seed_masters(session)
        first, second = _seed_order(session)
        _seed_batch(
            session,
            lines=[
                _line(first, 5, row_number=2),
                _line(first, -3, row_number=3),
                _line(
                    second,
                    2,
                    purchase_order_id="PO-2",
                    purchase_order_item_id="POI-2",
                    row_number=2,
                    sheet_name="来货测试工厂2",
                ),
            ],
        )

    result = _service(test_database_engine).confirm(
        batch_id=BATCH_ID, workbook_version=1, actor_id=ADMIN
    )

    assert result.record_count == 3
    assert result.factory_count == 2
    with Session(test_database_engine) as session:
        records = list(
            session.scalars(
                select(IncomingDiffRecord)
                .where(IncomingDiffRecord.order_id == ORDER_ID)
                .order_by(IncomingDiffRecord.registered_at, IncomingDiffRecord.record_id)
            )
        )
        assert sorted(record.quantity for record in records) == [-3, 2, 5]
        assert {record.initial_quantity for record in records} == {5, -3, 2}
        assert all(record.version == 1 for record in records)
        assert all(record.detail_id.startswith("idf-detail-") for record in records)
        assert all(record.product_name_snapshot == "来货测试产品" for record in records)
        assert all(record.product_code_snapshot == "IDF-ITEM" for record in records)
        assert all(record.source_business_date is None for record in records)
        audit = session.scalar(select(AuditLog).where(
            AuditLog.action == "incoming_diff_registered"
        ))
        assert audit is not None and "businessDate" not in audit.changes
        ledger = list(
            session.scalars(
                select(QuantityLedger)
                .where(QuantityLedger.source_type == "INCOMING_DIFF")
                .order_by(QuantityLedger.ledger_id)
            )
        )
        assert len(ledger) == 3
        assert {row.source_id for row in ledger} == {record.record_id for record in records}
        # AC-13: 多货可以让统一已发超过下单数量，本例 10 + 5 - 3 = 12
        assert _unified_shipped(session, first) == 12
        assert _unified_shipped(session, second) == 2
        batch = session.get(IncomingDiffBatch, BATCH_ID)
        assert batch is not None
        assert batch.status == "CONFIRMED"
        assert batch.confirmed_by == ADMIN
        assert batch.confirmed_record_count == 3
        assert batch.confirmed_factory_count == 2
        messages = list(
            session.scalars(
                select(OutboxMessage)
                .where(OutboxMessage.event_type == "incoming_diff.registered")
                .order_by(OutboxMessage.id)
            )
        )
        assert {message.dedupe_key for message in messages} == {
            f"incoming-diff:{BATCH_ID}:{FACTORY_A}:{ORDER_ID}",
            f"incoming-diff:{BATCH_ID}:{FACTORY_B}:{ORDER_ID}",
        }
        audits = list(
            session.scalars(
                select(AuditLog).where(AuditLog.action == "incoming_diff_registered")
            )
        )
        assert len(audits) == 1
        changes = audits[0].changes
        assert changes["recordCount"] == 3
        assert changes["assignments"][str(first)] == {"before": 10, "after": 12}
        assert changes["assignments"][str(second)] == {"before": 0, "after": 2}
        # 订单完成状态不因来货出入回退
        order = session.get(Order, ORDER_ID)
        assert order is not None and order.lifecycle == "PUBLISHED"


def test_confirm_allows_over_shipping_beyond_ordered_quantity(
    test_database_engine: Engine,
) -> None:
    with Session(test_database_engine) as session, session.begin():
        _seed_masters(session)
        first, _second = _seed_order(session)
        _seed_batch(session, lines=[_line(first, 200)])

    _service(test_database_engine).confirm(
        batch_id=BATCH_ID, workbook_version=1, actor_id=ADMIN
    )

    with Session(test_database_engine) as session:
        assignment_id = session.scalar(
            select(IncomingDiffRecord.order_assignment_id).where(
                IncomingDiffRecord.order_id == ORDER_ID
            )
        )
        assert assignment_id is not None
        assert _unified_shipped(session, assignment_id) == 210


def test_repeat_confirm_of_the_same_version_returns_the_original_result(
    test_database_engine: Engine,
) -> None:
    with Session(test_database_engine) as session, session.begin():
        _seed_masters(session)
        first, _second = _seed_order(session)
        _seed_batch(session, lines=[_line(first, 5)])

    service = _service(test_database_engine)
    first_result = service.confirm(batch_id=BATCH_ID, workbook_version=1, actor_id=ADMIN)
    second_result = service.confirm(batch_id=BATCH_ID, workbook_version=1, actor_id=ADMIN)

    assert second_result.record_count == first_result.record_count
    assert second_result.confirmed_at == first_result.confirmed_at
    assert second_result.replayed is True
    with Session(test_database_engine) as session:
        assert (
            session.scalar(
                select(func.count())
                .select_from(QuantityLedger)
                .where(QuantityLedger.source_type == "INCOMING_DIFF")
            )
            == 1
        )
        assert (
            session.scalar(
                select(func.count())
                .select_from(IncomingDiffRecord)
                .where(IncomingDiffRecord.order_id == ORDER_ID)
            )
            == 1
        )
        assert (
            session.scalar(
                select(func.count())
                .select_from(OutboxMessage)
                .where(OutboxMessage.event_type == "incoming_diff.registered")
            )
            == 1
        )


def test_one_failing_line_rolls_the_whole_batch_back(test_database_engine: Engine) -> None:
    with Session(test_database_engine) as session, session.begin():
        _seed_masters(session)
        first, second = _seed_order(session)
        _seed_batch(
            session,
            lines=[
                _line(first, 5, row_number=2),
                _line(second, 2, purchase_order_id=None, row_number=3),
            ],
        )

    with pytest.raises(IncomingDifferenceValidationError) as failure:
        _service(test_database_engine).confirm(
            batch_id=BATCH_ID, workbook_version=1, actor_id=ADMIN
        )

    assert failure.value.issues == [
        {"sheetName": "来货测试工厂1", "rowNumber": 3, "reason": "缺少采购主单号"}
    ]
    with Session(test_database_engine) as session:
        assert (
            session.scalar(select(func.count()).select_from(IncomingDiffRecord)) == 0
        )
        assert (
            session.scalar(
                select(func.count())
                .select_from(QuantityLedger)
                .where(QuantityLedger.source_type == "INCOMING_DIFF")
            )
            == 0
        )
        assert (
            session.scalar(
                select(func.count())
                .select_from(OutboxMessage)
                .where(OutboxMessage.event_type == "incoming_diff.registered")
            )
            == 0
        )
        batch = session.get(IncomingDiffBatch, BATCH_ID)
        assert batch is not None and batch.status == "READY"


def test_confirm_rejects_the_batch_when_unified_shipped_would_drop_below_zero(
    test_database_engine: Engine,
) -> None:
    with Session(test_database_engine) as session, session.begin():
        _seed_masters(session)
        first, second = _seed_order(session)
        _seed_batch(
            session,
            lines=[
                _line(first, 5, row_number=2),
                _line(
                    second,
                    -1,
                    purchase_order_id="PO-2",
                    purchase_order_item_id="POI-2",
                    row_number=3,
                    sheet_name="来货测试工厂2",
                ),
            ],
        )

    with pytest.raises(IncomingDifferenceValidationError) as failure:
        _service(test_database_engine).confirm(
            batch_id=BATCH_ID, workbook_version=1, actor_id=ADMIN
        )

    assert "统一已发数量" in str(failure.value)
    with Session(test_database_engine) as session:
        assert session.scalar(select(func.count()).select_from(IncomingDiffRecord)) == 0
        assert _unified_shipped(session, first) == 10
        assert _unified_shipped(session, second) == 0


def test_incoming_difference_does_not_net_against_a_receipt_difference(
    test_database_engine: Engine,
) -> None:
    with Session(test_database_engine) as session, session.begin():
        _seed_masters(session)
        first, _second = _seed_order(session)
        _seed_batch(session, lines=[_line(first, -2)])
        session.add(
            QuantityLedger(
                order_assignment_id=first,
                source_type="SHIPMENT_RECEIPT",
                source_id="idf-shipment",
                quantity_delta=-2,
                actor_id=ADMIN,
                created_at=SOURCE_TIME,
            )
        )

    _service(test_database_engine).confirm(
        batch_id=BATCH_ID, workbook_version=1, actor_id=ADMIN
    )

    with Session(test_database_engine) as session:
        assignment_id = session.scalar(
            select(IncomingDiffRecord.order_assignment_id).where(
                IncomingDiffRecord.order_id == ORDER_ID
            )
        )
        assert assignment_id is not None
        # 10 - 2（收货差额）- 2（来货出入）= 6，两个事实各自生效
        assert _unified_shipped(session, assignment_id) == 6
        assert (
            session.scalar(
                select(func.count())
                .select_from(QuantityLedger)
                .where(QuantityLedger.order_assignment_id == assignment_id)
            )
            == 2
        )


def test_adjustment_writes_delta_ledger_audit_and_outbox(
    test_database_engine: Engine,
) -> None:
    assignment_id, record_id = _confirmed_record(test_database_engine)

    result = _service(test_database_engine).adjust_quantity(
        actor_id=ADMIN,
        order_id=ORDER_ID,
        record_id=record_id,
        quantity=3,
        version=1,
        request_id="request-adjust-1",
    )

    assert result.quantity == 3
    assert result.version == 2
    with Session(test_database_engine) as session:
        assert _unified_shipped(session, assignment_id) == 13
        adjustment = session.scalar(
            select(IncomingDiffAdjustment).where(IncomingDiffAdjustment.record_id == record_id)
        )
        assert adjustment is not None
        assert (
            adjustment.before_quantity,
            adjustment.after_quantity,
            adjustment.delta,
            adjustment.request_id,
        ) == (5, 3, -2, "request-adjust-1")
        ledger = session.scalar(
            select(QuantityLedger).where(QuantityLedger.source_type == "INCOMING_DIFF_ADJUST")
        )
        assert ledger is not None
        assert ledger.source_id == adjustment.adjustment_id
        assert ledger.quantity_delta == -2
        audit = session.scalar(select(AuditLog).where(AuditLog.action == "incoming_diff_adjusted"))
        assert audit is not None
        assert audit.request_id == "request-adjust-1"
        assert audit.changes["beforeQuantity"] == 5
        assert audit.changes["afterQuantity"] == 3
        outbox = session.scalar(
            select(OutboxMessage).where(OutboxMessage.event_type == "incoming_diff.adjusted")
        )
        assert outbox is not None
        assert outbox.payload["factoryId"] == FACTORY_A
        assert outbox.payload["orderId"] == ORDER_ID


def test_noop_adjustment_returns_current_snapshot_without_writes(
    test_database_engine: Engine,
) -> None:
    _assignment_id, record_id = _confirmed_record(test_database_engine)

    result = _service(test_database_engine).adjust_quantity(
        actor_id=ADMIN,
        order_id=ORDER_ID,
        record_id=record_id,
        quantity=5,
        version=1,
        request_id="request-adjust-noop",
    )

    assert result.quantity == 5
    assert result.version == 1
    with Session(test_database_engine) as session:
        assert session.scalar(select(func.count()).select_from(IncomingDiffAdjustment)) == 0
        assert (
            session.scalar(
                select(func.count())
                .select_from(QuantityLedger)
                .where(QuantityLedger.source_type == "INCOMING_DIFF_ADJUST")
            )
            == 0
        )
        assert (
            session.scalar(
                select(func.count())
                .select_from(OutboxMessage)
                .where(OutboxMessage.event_type == "incoming_diff.adjusted")
            )
            == 0
        )
        assert (
            session.scalar(
                select(func.count())
                .select_from(AuditLog)
                .where(AuditLog.action == "incoming_diff_adjusted")
            )
            == 0
        )


def test_adjustment_rejects_stale_version_zero_and_negative_unified_quantity(
    test_database_engine: Engine,
) -> None:
    assignment_id, record_id = _confirmed_record(test_database_engine, quantity=-5)
    service = _service(test_database_engine)

    with pytest.raises(IncomingDifferenceConflict):
        service.adjust_quantity(
            actor_id=ADMIN,
            order_id=ORDER_ID,
            record_id=record_id,
            quantity=-6,
            version=2,
            request_id="request-adjust-stale",
        )
    with pytest.raises(IncomingDifferenceValidationError):
        service.adjust_quantity(
            actor_id=ADMIN,
            order_id=ORDER_ID,
            record_id=record_id,
            quantity=0,
            version=1,
            request_id="request-adjust-zero",
        )
    with pytest.raises(IncomingDifferenceValidationError):
        service.adjust_quantity(
            actor_id=ADMIN,
            order_id=ORDER_ID,
            record_id=record_id,
            quantity=-11,
            version=1,
            request_id="request-adjust-lower-bound",
        )

    with Session(test_database_engine) as session:
        record = session.get(IncomingDiffRecord, record_id)
        assert record is not None
        assert (record.quantity, record.version) == (-5, 1)
        assert _unified_shipped(session, assignment_id) == 5
        assert session.scalar(select(func.count()).select_from(IncomingDiffAdjustment)) == 0
        assert (
            session.scalar(
                select(func.count())
                .select_from(QuantityLedger)
                .where(QuantityLedger.source_type == "INCOMING_DIFF_ADJUST")
            )
            == 0
        )
        assert (
            session.scalar(
                select(func.count())
                .select_from(AuditLog)
                .where(AuditLog.action == "incoming_diff_adjusted")
            )
            == 0
        )
        assert (
            session.scalar(
                select(func.count())
                .select_from(OutboxMessage)
                .where(OutboxMessage.event_type == "incoming_diff.adjusted")
            )
            == 0
        )


def test_multiple_adjustments_accumulate_only_their_deltas(
    test_database_engine: Engine,
) -> None:
    assignment_id, record_id = _confirmed_record(test_database_engine)
    service = _service(test_database_engine)

    service.adjust_quantity(
        actor_id=ADMIN,
        order_id=ORDER_ID,
        record_id=record_id,
        quantity=3,
        version=1,
        request_id="request-adjust-first",
    )
    result = service.adjust_quantity(
        actor_id=ADMIN,
        order_id=ORDER_ID,
        record_id=record_id,
        quantity=8,
        version=2,
        request_id="request-adjust-second",
    )

    assert (result.quantity, result.version) == (8, 3)
    with Session(test_database_engine) as session:
        assert _unified_shipped(session, assignment_id) == 18
        assert (
            session.scalar(
                select(func.sum(IncomingDiffAdjustment.delta)).where(
                    IncomingDiffAdjustment.record_id == record_id
                )
            )
            == 3
        )
        assert (
            session.scalar(
                select(func.count())
                .select_from(IncomingDiffAdjustment)
                .where(IncomingDiffAdjustment.record_id == record_id)
            )
            == 2
        )


def test_adjustment_only_accepts_enabled_admin(
    test_database_engine: Engine,
) -> None:
    _assignment_id, record_id = _confirmed_record(test_database_engine)

    for actor_id in (FACTORY_USER, DISABLED_ADMIN):
        with pytest.raises(IncomingDifferencePermissionDenied):
            _service(test_database_engine).adjust_quantity(
                actor_id=actor_id,
                order_id=ORDER_ID,
                record_id=record_id,
                quantity=4,
                version=1,
                request_id=f"request-adjust-{actor_id}",
            )


def test_adjusted_event_notifies_every_enabled_factory_user(
    test_database_engine: Engine,
) -> None:
    _assignment_id, record_id = _confirmed_record(test_database_engine)
    with Session(test_database_engine) as session, session.begin():
        session.add(
            User(
                user_id=SECOND_FACTORY_USER,
                role="factory",
                is_enabled=True,
                feishu_display_name="同厂用户",
                factory_id=FACTORY_A,
                factory_position="employee",
            )
        )
        session.add(
            User(
                user_id=DISABLED_FACTORY_USER,
                role="factory",
                is_enabled=False,
                feishu_display_name="停用工厂用户",
                factory_id=FACTORY_A,
                factory_position="employee",
            )
        )
    _service(test_database_engine).adjust_quantity(
        actor_id=ADMIN,
        order_id=ORDER_ID,
        record_id=record_id,
        quantity=4,
        version=1,
        request_id="request-adjust-notify",
    )
    sessions = sessionmaker(test_database_engine, class_=Session, expire_on_commit=False)
    notifications = NotificationsAuditService(sessions)

    assert notifications.consume_next_business_event(worker_id="idf-worker") is True
    assert notifications.consume_next_business_event(worker_id="idf-worker") is True

    with Session(test_database_engine) as session:
        rows = list(
            session.scalars(
                select(Notification)
                .where(Notification.event_type == "incoming_diff.adjusted")
                .order_by(Notification.recipient_id)
            )
        )
        assert [row.recipient_id for row in rows] == [FACTORY_USER, SECOND_FACTORY_USER]
        assert {row.category for row in rows} == {"INCOMING_DIFF"}
        assert rows[0].title == "来货出入已调整"
        assert rows[0].target_path.startswith("/pages/factory-task-detail/factory-task-detail")


@pytest.mark.parametrize(
    ("actor_id", "submitter_id"),
    [
        (FACTORY_USER, ADMIN),
        (DISABLED_ADMIN, DISABLED_ADMIN),
        (OTHER_ADMIN, ADMIN),
    ],
)
def test_confirm_only_accepts_an_enabled_admin_confirming_their_own_batch(
    test_database_engine: Engine, actor_id: str, submitter_id: str
) -> None:
    with Session(test_database_engine) as session, session.begin():
        _seed_masters(session)
        first, _second = _seed_order(session)
        _seed_batch(session, lines=[_line(first, 5)], submitter_id=submitter_id)

    with pytest.raises(IncomingDifferencePermissionDenied):
        _service(test_database_engine).confirm(
            batch_id=BATCH_ID, workbook_version=1, actor_id=actor_id
        )

    with Session(test_database_engine) as session:
        assert session.scalar(select(func.count()).select_from(IncomingDiffRecord)) == 0
        assert (
            session.scalar(
                select(func.count())
                .select_from(QuantityLedger)
                .where(QuantityLedger.source_type == "INCOMING_DIFF")
            )
            == 0
        )


def test_batch_no_uses_shanghai_date_with_a_two_digit_daily_sequence(
    test_database_engine: Engine,
) -> None:
    with Session(test_database_engine) as session, session.begin():
        _seed_masters(session)
    # 2026-09-18 02:30 UTC 是上海时间 2026-09-18 10:30
    service = _service(test_database_engine)

    first = service.create_batch(submitter_id=ADMIN)
    second = service.create_batch(submitter_id=ADMIN)

    assert first.batch_no == "IN20260918-01"
    assert second.batch_no == "IN20260918-02"
    assert first.status == "COLLECTING"


def test_registered_event_notifies_every_enabled_factory_user(
    test_database_engine: Engine,
) -> None:
    with Session(test_database_engine) as session, session.begin():
        _seed_masters(session)
        first, _second = _seed_order(session)
        _seed_batch(session, lines=[_line(first, 5)])
        session.add(
            User(
                user_id=SECOND_FACTORY_USER,
                role="factory",
                is_enabled=True,
                feishu_display_name="同厂用户",
                factory_id=FACTORY_A,
                factory_position="employee",
            )
        )
        session.add(
            User(
                user_id=DISABLED_FACTORY_USER,
                role="factory",
                is_enabled=False,
                feishu_display_name="停用工厂用户",
                factory_id=FACTORY_A,
                factory_position="employee",
            )
        )
    _service(test_database_engine).confirm(
        batch_id=BATCH_ID, workbook_version=1, actor_id=ADMIN
    )
    sessions = sessionmaker(test_database_engine, class_=Session, expire_on_commit=False)
    notifications = NotificationsAuditService(sessions)

    assert notifications.consume_next_business_event(worker_id="idf-worker") is True

    with Session(test_database_engine) as session:
        rows = list(
            session.scalars(
                select(Notification)
                .where(Notification.event_type == "incoming_diff.registered")
                .order_by(Notification.recipient_id)
            )
        )
        assert [row.recipient_id for row in rows] == [FACTORY_USER, SECOND_FACTORY_USER]
        assert {row.category for row in rows} == {"INCOMING_DIFF"}
        assert rows[0].title == "订单来货出入已登记"
        assert "IDF-ORDER-1" in rows[0].summary
        assert rows[0].target_path.startswith("/pages/factory-task-detail/factory-task-detail")


def test_candidate_matching_requires_exact_sku_and_saved_purchase_link(
    test_database_engine: Engine,
) -> None:
    with Session(test_database_engine) as session, session.begin():
        _seed_masters(session)
        first, _second = _seed_order(session)
    service = _service(test_database_engine)

    assert (
        service.match_assignment(
            factory_id=FACTORY_A,
            product_code="IDF-ITEM",
            product_name="来货测试产品",
            spec="红色 / 110",
        )
        == first
    )

    # 缺少工厂或规格不能定位 SKU
    assert (
        service.match_assignment(
            product_code="IDF-ITEM",
            product_name="来货测试产品",
            spec="红色 / 110",
        )
        is None
    )
    assert (
        service.match_assignment(
            factory_id=FACTORY_A,
            product_code="IDF-ITEM",
            product_name="来货测试产品",
            spec="红色 / 110",
            purchase_order_id="PO-1",
        )
        == first
    )
    assert (
        service.match_assignment(
            factory_id=FACTORY_A,
            product_code="NOT-EXIST",
            product_name="来货测试产品",
            spec="红色 / 110",
        )
        is None
    )


def test_confirm_accepts_new_row_without_image_date(test_database_engine: Engine) -> None:
    with Session(test_database_engine) as session, session.begin():
        _seed_masters(session)
        first, _ = _seed_order(session)
        line = _line(first, -2)
        line["sourceBusinessDate"] = "legacy-invalid-date"
        _seed_batch(session, lines=[line])
    _service(test_database_engine).confirm(batch_id=BATCH_ID, workbook_version=1, actor_id=ADMIN)
    with Session(test_database_engine) as session:
        record = session.scalar(select(IncomingDiffRecord).where(
            IncomingDiffRecord.batch_id == BATCH_ID
        ))
        assert record is not None and record.source_business_date is None


def test_candidate_matching_skips_zero_unshipped_and_orders_by_contract_date(
    test_database_engine: Engine,
) -> None:
    with Session(test_database_engine) as session, session.begin():
        _seed_masters(session)
        first, _ = _seed_order(session)
        earliest = session.get(OrderAssignment, first)
        assert earliest is not None
        earliest.contract_ship_date = date(2026, 9, 10)
        session.add(Order(
            order_id="idf-order-2", order_no="IDF-ORDER-2", source="manual",
            order_date=date(2026, 9, 2), tracker="松子", trackers=["松子"],
            detail_mode=True, lifecycle="PUBLISHED", created_by=ADMIN,
            updated_by=ADMIN, created_at=SOURCE_TIME, updated_at=SOURCE_TIME,
        ))
        session.flush()
        line = OrderLine(
            order_id="idf-order-2", product_variant_id="idf-variant-1",
            order_quantity=20, sku_id_snapshot="IDF-SKU-1",
            product_name_snapshot="来货测试产品",
            properties_value_snapshot="红色 / 110", created_at=SOURCE_TIME,
            updated_at=SOURCE_TIME,
        )
        session.add(line)
        session.flush()
        later = OrderAssignment(
            order_line_id=line.order_line_id, factory_id=FACTORY_A, is_active=True,
            contract_ship_date=date(2026, 9, 20), assigned_quantity=20,
            initial_shipped_quantity=0, factory_name_snapshot="来货测试工厂1",
            created_at=SOURCE_TIME, updated_at=SOURCE_TIME,
        )
        session.add(later)
        session.flush()
        detail = OrderDetail(
            detail_id="idf-detail-later", order_id="idf-order-2", origin="manual",
            sort_order=1, accepted_raw_fields={}, purchase_order_id="PO-LATER",
            purchase_order_item_id="POI-LATER", source_sku_id="IDF-SKU-1",
            product_name="来货测试产品", properties_value="红色 / 110",
            matched_variant_id="idf-variant-1", matched_factory_id=FACTORY_A,
            order_quantity=20, source_trackers=["松子"],
            contract_ship_date=date(2026, 9, 20), parse_issues=[],
            assignment_id=later.order_assignment_id, dispatch_state="ASSIGNED",
            created_at=SOURCE_TIME, updated_at=SOURCE_TIME,
        )
        session.add(detail)
        session.flush()
        later.detail_id = detail.detail_id
        later_id = later.order_assignment_id
    service = _service(test_database_engine)
    query = {"factory_id": FACTORY_A, "product_code": "IDF-ITEM", "spec": "红色 / 110"}
    assert service.match_assignment(**query) == first
    with Session(test_database_engine) as session, session.begin():
        session.add(QuantityLedger(
            order_assignment_id=first, source_type="ADMIN_ADJUST", source_id="idf-ledger-1",
            quantity_delta=90, actor_id=ADMIN, created_at=SOURCE_TIME,
        ))
    assert service.match_assignment(**query) == later_id
    with Session(test_database_engine) as session, session.begin():
        later = session.get(OrderAssignment, later_id)
        assert later is not None
        later.initial_shipped_quantity = later.assigned_quantity
    assert service.match_assignment(**query) is None


def test_recognition_worker_retries_and_only_saves_candidates(
    test_database_engine: Engine,
) -> None:
    content = b"private-fake-image"
    digest = hashlib.sha256(content).hexdigest()
    files = FakePrivateFileStore(bucket="incoming-test")
    files.put(object_key="images/vision.jpg", content=content, content_type="image/jpeg")
    with Session(test_database_engine) as session, session.begin():
        _seed_masters(session)
        first, _ = _seed_order(session)
        session.add(StoredFile(
            file_id=9709, bucket=files.bucket, object_key="images/vision.jpg",
            original_filename="vision.jpg", mime_type="image/jpeg",
            size_bytes=len(content), content_sha256=digest, uploaded_by=ADMIN,
        ))
    sessions = sessionmaker(test_database_engine, class_=Session, expire_on_commit=False)
    payload = json.dumps({
        "factoryName": "来货测试工厂1", "productCode": "IDF-ITEM",
        "productName": "来货测试产品", "boxes": [{"color": "红色", "size": "110"}],
        "diffNoteText": "少2，多3", "lines": [
            {"color": "红色", "size": "110", "direction": "少", "quantity": 2},
            {"color": "红色", "size": "110", "direction": "多", "quantity": 3},
        ], "confidence": 0.8, "unresolvedFields": [],
    }, ensure_ascii=False)
    fake = FakeIncomingDiffRecognizer([ValueError("temporary"), payload])
    recognition = IncomingDiffRecognitionService(sessions, files=files, recognizer=fake,
                                                 clock=lambda: NOW)
    batch = _service(test_database_engine).create_batch(submitter_id=ADMIN)
    image = recognition.add_image(batch_id=batch.batch_id, actor_id=ADMIN,
                                  file_id=9709, feishu_image_key="vision-key")
    job_id = recognition.freeze_and_enqueue(batch_id=batch.batch_id, actor_id=ADMIN)
    store = InfrastructureStore(sessions)
    worker = Worker(store=store, worker_id="vision-test",
                    handlers={"incoming_diff.recognize": recognition.recognize},
                    terminal_failure_handlers={
                        "incoming_diff.recognize": recognition.fail_terminal},
                    retry_limits={"incoming_diff.recognize": 3}, retry_delay_seconds=0)
    now = NOW.replace(tzinfo=None)
    assert worker.run_once(now=now)
    assert store.get_job(job_id=job_id).status == "pending"
    assert worker.run_once(now=now)
    assert store.get_job(job_id=job_id).status == "completed"
    with Session(test_database_engine) as session:
        saved = session.get(IncomingDiffImage, image.image_id)
        assert saved is not None and saved.ocr_status == "SUCCEEDED"
        assert saved.recognition_model == "qwen3.5-flash"
        assert [line["signedQuantity"] for line in saved.recognition_payload["lines"]] == [-2, 3]
        assert all(line["orderAssignmentId"] == first
                   for line in saved.recognition_payload["lines"])
        assert session.scalar(select(func.count()).select_from(IncomingDiffRecord)) == 0
        assert session.scalar(select(func.count()).select_from(QuantityLedger)) == 0


def test_duplicate_image_needs_explicit_acknowledgement(test_database_engine: Engine) -> None:
    with Session(test_database_engine) as session, session.begin():
        _seed_masters(session)
        session.add(IncomingDiffBatch(
            batch_id="idf-confirmed", batch_no="IN20260917-01", submitter_id=ADMIN,
            status="CONFIRMED", confirmed_at=SOURCE_TIME, created_at=SOURCE_TIME,
            updated_at=SOURCE_TIME,
        ))
        session.flush()
        session.add(IncomingDiffImage(
            image_id="idf-old-image", batch_id="idf-confirmed", sort_order=1,
            feishu_image_key="old-key", file_id=FILE_IDS[0],
            content_sha256="a" * 64, ocr_status="SUCCEEDED", created_at=SOURCE_TIME,
        ))
    sessions = sessionmaker(test_database_engine, class_=Session, expire_on_commit=False)
    recognition = IncomingDiffRecognitionService(
        sessions, files=FakePrivateFileStore(bucket="incoming-test"),
        recognizer=FakeIncomingDiffRecognizer([]), clock=lambda: NOW,
    )
    batch = _service(test_database_engine).create_batch(submitter_id=ADMIN)
    image = recognition.add_image(batch_id=batch.batch_id, actor_id=ADMIN,
                                  file_id=FILE_IDS[0], feishu_image_key="new-key")
    assert image.duplicate_of_batch_id == "idf-confirmed"
    with pytest.raises(ValueError, match="重复图片"):
        recognition.freeze_and_enqueue(batch_id=batch.batch_id, actor_id=ADMIN)
    recognition.acknowledge_duplicate(batch_id=batch.batch_id, image_id=image.image_id,
                                      actor_id=ADMIN)
    recognition.freeze_and_enqueue(batch_id=batch.batch_id, actor_id=ADMIN)


def test_same_batch_recognizes_different_products_per_image(test_database_engine: Engine) -> None:
    files = FakePrivateFileStore(bucket="incoming-test")
    with Session(test_database_engine) as session, session.begin():
        _seed_masters(session)
        _seed_order(session)
        for file_id, content in ((9709, b"product-one"), (9710, b"product-two")):
            key = f"images/{file_id}.jpg"
            files.put(object_key=key, content=content, content_type="image/jpeg")
            session.add(StoredFile(
                file_id=file_id, bucket=files.bucket, object_key=key,
                original_filename=f"{file_id}.jpg", mime_type="image/jpeg",
                size_bytes=len(content), content_sha256=hashlib.sha256(content).hexdigest(),
                uploaded_by=ADMIN,
            ))
    responses = [json.dumps({
        "factoryName": "来货测试工厂1", "productCode": code,
        "productName": name, "boxes": [{"color": "红色", "size": "110"}],
        "diffNoteText": "少1", "lines": [{"color": "红色", "size": "110",
                                  "direction": "少", "quantity": 1}],
        "confidence": 0.9, "unresolvedFields": [],
    }, ensure_ascii=False) for code, name in (
        ("IDF-ITEM", "来货测试产品"), ("OTHER-ITEM", "另一产品")
    )]
    fake = FakeIncomingDiffRecognizer(responses)
    sessions = sessionmaker(test_database_engine, class_=Session, expire_on_commit=False)
    recognition = IncomingDiffRecognitionService(sessions, files=files, recognizer=fake,
                                                 clock=lambda: NOW)
    batch = _service(test_database_engine).create_batch(submitter_id=ADMIN)
    for file_id in (9709, 9710):
        recognition.add_image(batch_id=batch.batch_id, actor_id=ADMIN,
                              file_id=file_id, feishu_image_key=f"key-{file_id}")
    recognition.freeze_and_enqueue(batch_id=batch.batch_id, actor_id=ADMIN)
    worker = Worker(store=InfrastructureStore(sessions), worker_id="vision-test",
                    handlers={"incoming_diff.recognize": recognition.recognize})
    assert worker.run_once(now=NOW.replace(tzinfo=None))
    assert fake.calls == 2
    with Session(test_database_engine) as session:
        images = session.scalars(select(IncomingDiffImage).where(
            IncomingDiffImage.batch_id == batch.batch_id
        ).order_by(IncomingDiffImage.sort_order)).all()
        assert [image.recognition_payload["productCode"] for image in images] == [
            "IDF-ITEM", "OTHER-ITEM"
        ]
        assert images[0].recognition_payload["lines"][0]["orderAssignmentId"] is not None
        assert images[1].recognition_payload["lines"][0]["orderAssignmentId"] is None


def test_recognition_reaches_terminal_failure(test_database_engine: Engine) -> None:
    content = b"private-fake-image"
    digest = hashlib.sha256(content).hexdigest()
    files = FakePrivateFileStore(bucket="incoming-test")
    files.put(object_key="images/vision.jpg", content=content, content_type="image/jpeg")
    with Session(test_database_engine) as session, session.begin():
        _seed_masters(session)
        session.add(StoredFile(
            file_id=9709, bucket=files.bucket, object_key="images/vision.jpg",
            original_filename="vision.jpg", mime_type="image/jpeg",
            size_bytes=len(content), content_sha256=digest, uploaded_by=ADMIN,
        ))
    sessions = sessionmaker(test_database_engine, class_=Session, expire_on_commit=False)
    fake = FakeIncomingDiffRecognizer(["{", "{", "{"])
    recognition = IncomingDiffRecognitionService(sessions, files=files, recognizer=fake,
                                                 clock=lambda: NOW)
    batch = _service(test_database_engine).create_batch(submitter_id=ADMIN)
    image = recognition.add_image(batch_id=batch.batch_id, actor_id=ADMIN,
                                  file_id=9709, feishu_image_key="vision-key")
    recognition.freeze_and_enqueue(batch_id=batch.batch_id, actor_id=ADMIN)
    worker = Worker(store=InfrastructureStore(sessions), worker_id="vision-test",
                    handlers={"incoming_diff.recognize": recognition.recognize},
                    terminal_failure_handlers={
                        "incoming_diff.recognize": recognition.fail_terminal},
                    retry_limits={"incoming_diff.recognize": 3}, retry_delay_seconds=0)
    now = NOW.replace(tzinfo=None)
    for _ in range(3):
        assert worker.run_once(now=now)
    with Session(test_database_engine) as session:
        failed = session.get(IncomingDiffImage, image.image_id)
        failed_batch = session.get(IncomingDiffBatch, batch.batch_id)
        assert failed is not None and failed.ocr_status == "FAILED"
        assert failed.failure_reason == "invalid_json"
        assert failed_batch is not None and failed_batch.status == "FAILED"
        assert failed_batch.recognition_error_summary == "第 1 张图片识别失败"


class _BotMedia:
    def __init__(self, image: bytes) -> None:
        self.image = image
        self.workbook = b""
        self.downloads: list[str] = []

    def download_resource(self, message_id: str, file_key: str, resource_type: str) -> bytes:
        self.downloads.append(resource_type)
        return self.image if resource_type == "image" else self.workbook


def _bot_event(event_id: str, *, kind: str = "image", open_id: str = "open-1",
               filename: str = "") -> dict[str, object]:
    content = {"image_key": "image-key"} if kind == "image" else {
        "file_key": "workbook-key", "file_name": filename,
    }
    return {
        "schema": "2.0", "header": {"event_id": event_id,
                                    "event_type": "im.message.receive_v1",
                                    "tenant_key": "tenant-test"},
        "event": {"sender": {"sender_id": {"open_id": open_id}},
                  "message": {"message_id": f"msg-{event_id}", "chat_id": "chat-1",
                              "chat_type": "p2p", "message_type": kind,
                              "content": json.dumps(content)}},
    }


def _bot_action(event_id: str, batch_id: str, action: str,
                *, version: int | None = None, open_id: str = "open-1") -> dict[str, object]:
    value: dict[str, object] = {"batchId": batch_id, "action": action}
    if version is not None:
        value["version"] = version
    return {"schema": "2.0", "header": {"event_id": event_id,
                                        "event_type": "card.action.trigger",
                                        "tenant_key": "tenant-test"},
            "event": {"operator": {"open_id": open_id}, "action": {"value": value}}}


def test_feishu_bot_photo_workbook_upload_confirm_and_replay(
    test_database_engine: Engine,
) -> None:
    with Session(test_database_engine) as session, session.begin():
        _seed_masters(session)
        assignment_id, _ = _seed_order(session)
        session.add_all([
            ExternalIdentity(platform="feishu", scope="test-scope",
                             platform_subject=f"tenant-test:{open_id}", user_id=user_id)
            for open_id, user_id in (("open-1", ADMIN), ("open-2", OTHER_ADMIN))
        ])
    picture = BytesIO()
    Image.new("RGB", (4, 4), (255, 255, 255)).save(picture, format="PNG")
    media = _BotMedia(picture.getvalue())
    files = FakePrivateFileStore(bucket="incoming-test")
    sessions = sessionmaker(test_database_engine, class_=Session, expire_on_commit=False)
    response = json.dumps({
        "factoryName": "来货测试工厂1", "productCode": "IDF-ITEM",
        "productName": "来货测试产品", "boxes": [{"color": "红色", "size": "110"}],
        "diffNoteText": "多2", "lines": [
            {"color": "红色", "size": "110", "direction": "多", "quantity": 2}
        ], "confidence": 0.9, "unresolvedFields": [],
    }, ensure_ascii=False)
    recognizer = FakeIncomingDiffRecognizer([response])
    recognition = IncomingDiffRecognitionService(
        sessions, files=files, recognizer=recognizer, clock=lambda: NOW
    )
    bot = FeishuBotService(
        sessions, files=files, media=media, identity_scope="test-scope",
        codec=IncomingWorkbookCodec(Settings(database_url="mysql+pymysql://local/test")),
        recognition=recognition,
    )
    assert bot.event(_bot_event("unrelated", kind="text")) == {}
    bot.event(_bot_event("denied", open_id="unknown"))
    with Session(test_database_engine) as session:
        assert session.scalar(select(func.count()).select_from(IncomingDiffBatch)) == 0
    bot.event(_bot_event("image-1"))
    bot.event(_bot_event("image-1"))
    bot.release_failed_event(_bot_event("image-1"))
    bot.event(_bot_event("image-1"))
    with Session(test_database_engine) as session:
        batch = session.scalar(select(IncomingDiffBatch))
        assert batch is not None
        batch_id, batch_no = batch.batch_id, batch.batch_no
        assert session.scalar(select(func.count()).select_from(IncomingDiffImage)) == 1
    bot.card_action(_bot_action("other-confirm", batch_id, "generate", open_id="open-2"))
    flat_action = _bot_action("flat-other", batch_id, "generate", open_id="open-2")
    flat_action = {"schema": "2.0", **flat_action["header"], **flat_action["event"]}
    flat_action.pop("tenant_key")
    flat_action["operator"]["tenant_key"] = "tenant-test"
    assert bot.card_action(flat_action)["toast"]["content"] == "只能操作本人批次"
    with Session(test_database_engine) as session:
        assert session.get(IncomingDiffBatch, batch_id).status == "COLLECTING"
    bot.card_action(_bot_action("generate", batch_id, "generate"))
    worker = Worker(
        store=InfrastructureStore(sessions), worker_id="bot-test",
        handlers={"incoming_diff.recognize": bot.recognition_job,
                  CONFIRM_JOB: bot.confirm_job},
    )
    assert worker.run_once()
    with Session(test_database_engine) as session:
        batch = session.get(IncomingDiffBatch, batch_id)
        assert batch is not None and batch.status == "READY"
        workbook = session.get(IncomingDiffWorkbook, batch.current_workbook_id)
        assert workbook is not None and workbook.version == 1
        stored = session.get(StoredFile, workbook.file_id)
        assert stored is not None
        edited = load_workbook(BytesIO(files.get(object_key=stored.object_key)))
        edited["来货测试工厂1"]["C2"] = 3
        output = BytesIO()
        edited.save(output)
        media.workbook = output.getvalue()
        assert workbook.line_snapshot[0]["orderAssignmentId"] == assignment_id
    bot.event(_bot_event("upload", kind="file", filename=f"{batch_no}_核对表_v1.xlsx"))
    assert media.downloads == ["image", "file"]
    with Session(test_database_engine) as session:
        batch = session.get(IncomingDiffBatch, batch_id)
        uploaded = session.get(IncomingDiffWorkbook, batch.current_workbook_id)
        assert uploaded is not None and uploaded.version == 2
    bot.card_action(_bot_action("confirm", batch_id, "confirm", version=2))
    assert worker.run_once()
    bot.card_action(_bot_action("confirm-again", batch_id, "confirm", version=2))
    with Session(test_database_engine) as session:
        assert session.get(IncomingDiffBatch, batch_id).status == "CONFIRMED"
        assert session.scalar(select(func.count()).select_from(IncomingDiffRecord)) == 1
        delta = session.scalar(select(func.sum(QuantityLedger.quantity_delta)).where(
            QuantityLedger.order_assignment_id == assignment_id
        ))
        assert delta == 3
        assert session.scalar(select(func.count()).select_from(OutboxMessage).where(
            OutboxMessage.event_type == "incoming_diff.bot_reply"
        )) >= 5

    class FailingOnceNotifier:
        def __init__(self) -> None:
            self.sent: list[DeliveryRequest] = []
            self.failed = False

        def send(self, request: DeliveryRequest) -> None:
            if "已正式登记" in request.summary and not self.failed:
                self.failed = True
                raise NotificationDeliveryError("temporary", retryable=True)
            self.sent.append(request)

    notifier = FailingOnceNotifier()
    notifications = NotificationsAuditService(sessions)
    while notifications.deliver_next(
        worker_id="bot-delivery", wechat_notifier=FakeWechatNotifier(),
        feishu_notifier=notifier, enabled_channels={"feishu"},
    ):
        pass
    assert notifier.failed
    assert any(request.file_id == workbook.file_id for request in notifier.sent)
    with Session(test_database_engine) as session:
        assert session.scalar(select(func.sum(QuantityLedger.quantity_delta)).where(
            QuantityLedger.order_assignment_id == assignment_id
        )) == 3
    assert notifications.deliver_next(
        worker_id="bot-delivery", wechat_notifier=FakeWechatNotifier(),
        feishu_notifier=notifier, enabled_channels={"feishu"},
        now=utc_now() + timedelta(seconds=31),
    )
    assert any("已正式登记" in request.summary for request in notifier.sent)


def test_feishu_bot_failed_download_marks_image_and_starts_new_batch(
    test_database_engine: Engine,
) -> None:
    with Session(test_database_engine) as session, session.begin():
        _seed_masters(session)
        session.add(ExternalIdentity(
            platform="feishu", scope="test-scope",
            platform_subject="tenant-test:open-1", user_id=ADMIN,
        ))
    media = _BotMedia(b"")
    files = FakePrivateFileStore(bucket="incoming-test")
    sessions = sessionmaker(test_database_engine, class_=Session, expire_on_commit=False)
    recognition = IncomingDiffRecognitionService(
        sessions, files=files, recognizer=FakeIncomingDiffRecognizer([]), clock=lambda: NOW
    )
    bot = FeishuBotService(
        sessions, files=files, media=media, identity_scope="test-scope",
        codec=IncomingWorkbookCodec(Settings(database_url="mysql+pymysql://local/test")),
        recognition=recognition,
    )
    bot.event(_bot_event("failed-image"))
    with Session(test_database_engine) as session:
        first = session.scalar(select(IncomingDiffBatch))
        image = session.scalar(select(IncomingDiffImage))
        assert first is not None and first.status == "FAILED"
        assert image is not None and image.ocr_status == "FAILED" and image.file_id is None
        assert files.object_count == 0
    picture = BytesIO()
    Image.new("RGB", (4, 4), (255, 255, 255)).save(picture, format="PNG")
    media.image = picture.getvalue()
    bot.event(_bot_event("new-image"))
    with Session(test_database_engine) as session:
        assert session.scalar(select(func.count()).select_from(IncomingDiffBatch)) == 2
        assert session.scalar(select(func.count()).select_from(IncomingDiffImage)) == 2


def test_feishu_bot_reports_ambiguous_purchase_suborders(
    test_database_engine: Engine,
) -> None:
    with Session(test_database_engine) as session, session.begin():
        _seed_masters(session)
        _first, second = _seed_order(session)
        session.add(Order(
            order_id="idf-order-2", order_no="IDF-ORDER-2", source="manual",
            order_date=date(2026, 9, 2), tracker="松子", trackers=["松子"],
            detail_mode=True, lifecycle="PUBLISHED", created_by=ADMIN,
            updated_by=ADMIN, created_at=SOURCE_TIME, updated_at=SOURCE_TIME,
        ))
        session.flush()
        assignment = session.get(OrderAssignment, second)
        assert assignment is not None
        assignment.factory_id = FACTORY_A
        order_line = session.get(OrderLine, assignment.order_line_id)
        assert order_line is not None
        order_line.order_id = "idf-order-2"
        order_line.product_variant_id = "idf-variant-1"
        detail = session.get(OrderDetail, assignment.detail_id)
        assert detail is not None
        detail.order_id = "idf-order-2"
        detail.purchase_order_id = "PO-1"
        detail.matched_variant_id = "idf-variant-1"
        detail.matched_factory_id = FACTORY_A
    sessions = sessionmaker(test_database_engine, class_=Session, expire_on_commit=False)
    files = FakePrivateFileStore(bucket="incoming-test")
    bot = FeishuBotService(
        sessions, files=files, media=_BotMedia(b""), identity_scope="test-scope",
        codec=IncomingWorkbookCodec(Settings(database_url="mysql+pymysql://local/test")),
        recognition=IncomingDiffRecognitionService(
            sessions, files=files, recognizer=FakeIncomingDiffRecognizer([])),
    )
    message = bot._registration_errors(
        [{"sheetName": "来货测试工厂1", "rowNumber": 2,
          "reason": "未匹配到唯一派工"}],
        [{"sheetName": "来货测试工厂1", "rowNumber": 2,
          "productCode": "IDF-ITEM", "productName": "来货测试产品",
          "spec": "红色 / 110", "purchaseOrderId": "PO-1"}],
    )
    assert message == (
        "来货测试工厂1 第 2 行的采购单号 PO-1 存在 2 个符合条件的子单，"
        "无法确定归属。请补充或修正后重新发送。"
    )
    with Session(test_database_engine) as session, session.begin():
        assignment = session.get(OrderAssignment, second)
        assert assignment is not None
        assignment.detail_id = None


def test_factory_notification_failure_queues_bot_notice_without_reconfirming(
    test_database_engine: Engine,
) -> None:
    assignment_id, _ = _confirmed_record(test_database_engine, quantity=2)
    with Session(test_database_engine) as session, session.begin():
        session.add(NotificationAuthorization(
            user_id=FACTORY_USER, template_key="factory_status", result="accepted",
            authorized_at=SOURCE_TIME,
        ))
    sessions = sessionmaker(test_database_engine, class_=Session, expire_on_commit=False)
    service = NotificationsAuditService(sessions)
    assert service.consume_next_business_event(worker_id="factory-notice")

    class FailingWechat:
        def send(self, request: DeliveryRequest) -> None:
            raise NotificationDeliveryError("temporary", retryable=True)

    assert service.deliver_next(
        worker_id="factory-notice", wechat_notifier=FailingWechat(),
        enabled_channels={"wechat"},
    )
    with Session(test_database_engine) as session:
        bot_notice = session.scalar(select(OutboxMessage).where(
            OutboxMessage.dedupe_key.like("incoming-diff-bot:notify-failed:%")
        ))
        assert bot_notice is not None
        assert "数量不会重复调整" in bot_notice.payload["summary"]
        assert session.scalar(select(func.sum(QuantityLedger.quantity_delta)).where(
            QuantityLedger.order_assignment_id == assignment_id,
            QuantityLedger.source_type == "INCOMING_DIFF",
        )) == 2
    assert service.deliver_next(
        worker_id="factory-notice", wechat_notifier=FakeWechatNotifier(),
        enabled_channels={"wechat"}, now=utc_now() + timedelta(seconds=31),
    )
    with Session(test_database_engine) as session:
        assert session.scalar(select(func.count()).select_from(IncomingDiffRecord)) == 1
