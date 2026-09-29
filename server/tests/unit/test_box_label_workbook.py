from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook
from PIL import Image

from app.modules.box_labels.service import BoxLabelService, group_id
from app.modules.box_labels.workbook import BoxLabelWorkbookRenderer
from app.modules.product_sync.color import extract_color

TEMPLATE = Path(__file__).resolve().parents[2] / "app/templates/box_label_v1.xlsx"


def test_real_spec_formats_and_ambiguous_color() -> None:
    cases = {
        "冰川蓝110": "冰川蓝",
        "红色120/60": "红色",
        "粉色;XL": "粉色",
        "咖啡;": "咖啡",
        "米色 / 52cm": "米色",
        "23号冰川蓝均码": "23号冰川蓝",
        "23号冰川蓝无尺码": "23号冰川蓝",
        "蓝色小熊;S": "蓝色小熊",
        "冰川蓝": "冰川蓝",
        "蓝2": None,
        "冰川蓝2026": None,
        "": None,
    }
    assert {value: extract_color(value) for value in cases} == cases
    assert group_id("factory", "product-1", "蓝色") != group_id(
        "factory", "product-2", "蓝色"
    )
    assert BoxLabelService._filename("092#/", {
        "factory": "工/厂", "itemNo": "KQ26168",
        "productName": "暖呼吸", "color": "冰川蓝",
    }) == "092#__工_厂_暖呼吸_冰川蓝.xlsx"


def test_template_and_rendered_labels_keep_two_printable_copies() -> None:
    template = load_workbook(TEMPLATE)
    sheet = template.active
    assert sheet is not None
    assert len(sheet._images) == 0
    assert str(sheet.print_area).endswith("$A$1:$E$18")
    assert sheet["B1"].value == "工厂："
    assert sheet["B11"].value == "工厂："
    assert sheet["B5"].value == sheet["B15"].value == "尺码："
    assert sheet["B8"].value == sheet["B18"].value == "箱号/总箱数："
    assert sheet["E18"].border.right.style is not None

    source = BytesIO()
    Image.new("RGB", (200, 100), "red").save(source, format="PNG")
    values = {
        "factory": "晟衣", "itemNo": "KQ26168",
        "productName": "暖呼吸羽绒马甲", "color": "冰川蓝",
    }
    content = BoxLabelWorkbookRenderer().render(values, source.getvalue())
    rendered = load_workbook(BytesIO(content))
    result = rendered.active
    assert result is not None
    assert len(result._images) == 0
    # openpyxl 不识别单元格图片；WPS 实际显示另行验收。
    with ZipFile(BytesIO(content)) as package:
        assert "xl/richData/rdrichvalue.xml" in package.namelist()
        assert "xl/richData/_rels/richValueRel.xml.rels" in package.namelist()
        assert "/xl/richData/_rels/richValueRel.xml.rels" not in package.namelist()
        assert "xl/media/image1.png" in package.namelist()
    assert result["A2"].value == result["A12"].value == "#VALUE!"
    assert str(result.print_area).endswith("$A$1:$E$18")
    assert {str(item) for item in result.merged_cells.ranges} == {
        str(item) for item in sheet.merged_cells.ranges
    }
    assert result.sheet_format.defaultRowHeight == sheet.sheet_format.defaultRowHeight
    for start in (1, 11):
        image_cell = result.cell(start + 1, 1)
        template_cell = sheet.cell(start + 1, 1)
        assert image_cell.border.left.style == template_cell.border.left.style
        assert image_cell.border.top.style == template_cell.border.top.style
        assert result.cell(start, 2).value == "工厂：晟衣"
        assert result.cell(start + 1, 2).value == "货号：KQ26168"
        assert result.cell(start + 2, 2).value == "品名：暖呼吸羽绒马甲"
        assert result.cell(start + 3, 2).value == "颜色：冰川蓝"
        for offset in (4, 5, 6, 7):
            assert result.cell(start + offset, 2).value == sheet.cell(start + offset, 2).value


def test_webp_image_is_embedded_as_png() -> None:
    source = BytesIO()
    Image.new("RGB", (20, 10), "red").save(source, format="WEBP")
    values = {
        "factory": "晟衣", "itemNo": "KQ26168",
        "productName": "暖呼吸羽绒马甲", "color": "冰川蓝",
    }
    content = BoxLabelWorkbookRenderer().render(values, source.getvalue())
    with ZipFile(BytesIO(content)) as package:
        assert "xl/media/image1.png" in package.namelist()
