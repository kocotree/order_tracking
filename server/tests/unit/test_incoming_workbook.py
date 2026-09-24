from datetime import UTC, datetime
from io import BytesIO

import pytest
from openpyxl import load_workbook
from PIL import Image

from app.modules.incoming_differences.workbook import (
    HEADERS,
    IncomingWorkbookCodec,
    IncomingWorkbookLimits,
    IncomingWorkbookValidationError,
)
from app.settings.config import Settings

BATCH = "batch-A"
IMAGE = "image-A"


def _codec(*, limits: IncomingWorkbookLimits | None = None) -> IncomingWorkbookCodec:
    return IncomingWorkbookCodec(
        Settings(database_url="mysql+pymysql://fake:fake@localhost/fake"), limits=limits
    )


def _image() -> bytes:
    output = BytesIO()
    Image.new("RGB", (1200, 600), "red").save(output, format="PNG")
    return output.getvalue()


def _source():
    lines = [
        {
            "imageId": IMAGE, "factoryName": "甲工厂", "productName": "帽子",
            "spec": "红 / 110", "quantity": -1, "purchaseOrderId": "PO-1",
        },
        {
            "imageId": IMAGE, "factoryName": "甲工厂", "productName": "帽子",
            "spec": "红 / 120", "quantity": 2,
        },
        {
            "imageId": IMAGE, "factoryName": None, "productName": "帽子",
            "spec": "蓝 / 130", "quantity": 1,
        },
    ]
    return _codec().generate(
        batch_id=BATCH, version=1, lines=lines, images={IMAGE: _image()},
        generated_at=datetime(2026, 9, 24, tzinfo=UTC),
    )


def _edit(content: bytes, change) -> bytes:
    workbook = load_workbook(BytesIO(content))
    change(workbook)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _codes(error: pytest.ExceptionInfo[IncomingWorkbookValidationError]) -> list[str]:
    return [str(issue["code"]) for issue in error.value.issues]


def test_deployment_rejects_local_fake_signing_key() -> None:
    settings = Settings(database_url="mysql+pymysql://fake:fake@localhost/fake")
    with pytest.raises(ValueError, match="签名密钥未配置"):
        IncomingWorkbookCodec(settings.model_copy(update={"app_env": "production"}))
    assert settings.incoming_diff_workbook_signing_secret not in repr(settings)


def test_generate_structure_and_ac07_edits() -> None:
    content, signature, original = _source()
    workbook = load_workbook(BytesIO(content))
    assert workbook.sheetnames == ["甲工厂", "待确认", "_核对标识"]
    assert workbook["_核对标识"].sheet_state == "hidden"
    assert workbook["甲工厂"].column_dimensions["I"].hidden
    assert tuple(workbook["甲工厂"].cell(1, col).value for col in range(1, 9)) == HEADERS
    assert workbook["甲工厂"].row_dimensions[2].height == 96
    assert len(workbook["甲工厂"]._images) == 2
    assert workbook["甲工厂"]["D2"].value is None
    assert workbook["甲工厂"]["E2"].value is None
    assert workbook["甲工厂"]["G2"].value is None

    def change(book):
        first = book["甲工厂"]
        first["C2"] = -3
        moved = [first.cell(3, col).value for col in range(1, 10)]
        first.delete_rows(3)
        book["待确认"].append(moved)
        book["待确认"].delete_rows(2)

    edited = _edit(content, change)
    parsed = _codec().parse(
        edited, batch_id=BATCH, version=1, signature=signature, previous_lines=original
    )
    observed = [
        (line["sheetName"], line["rowNumber"], line["quantity"], line["imageId"])
        for line in parsed
    ]
    assert observed == [
        ("甲工厂", 2, -3, IMAGE), ("待确认", 2, 2, IMAGE)
    ]
    assert parsed[1]["factoryName"] is None
    advanced, next_signature = _codec().advance(
        edited, batch_id=BATCH, version=2, lines=parsed,
        generated_at=datetime(2026, 9, 25, tzinfo=UTC),
    )
    assert _codec().parse(
        advanced, batch_id=BATCH, version=2, signature=next_signature,
        previous_lines=parsed,
    ) == parsed


def test_missing_token_reports_sheet_and_row() -> None:
    content, signature, lines = _source()
    edited = _edit(content, lambda book: book["甲工厂"].append(["新增", "蓝", 1]))
    with pytest.raises(IncomingWorkbookValidationError) as caught:
        _codec().parse(edited, batch_id=BATCH, version=1, signature=signature, previous_lines=lines)
    assert caught.value.issues == [{
        "code": "missing_token", "sheet": "甲工厂", "row": 4,
        "message": "数据行缺少行标识",
    }]


def test_token_from_another_batch_is_rejected() -> None:
    content, signature, lines = _source()
    _, _, other_lines = _codec().generate(
        batch_id="batch-B", version=1,
        lines=[{
            "imageId": IMAGE, "factoryName": "甲工厂", "productName": "帽子",
            "spec": "红 / 110", "quantity": 1,
        }],
        images={IMAGE: _image()}, generated_at=datetime(2026, 9, 24, tzinfo=UTC),
    )
    edited = _edit(
        content,
        lambda book: book["甲工厂"].__setitem__("I2", other_lines[0]["lineToken"]),
    )
    with pytest.raises(IncomingWorkbookValidationError) as caught:
        _codec().parse(
            edited, batch_id=BATCH, version=1,
            signature=signature, previous_lines=lines,
        )
    assert caught.value.issues[0]["code"] == "foreign_token"


@pytest.mark.parametrize(
    ("change", "batch", "version", "code"),
    [
        (
            lambda book: book["_核对标识"].__setitem__("B4", "tampered"),
            BATCH, 1, "invalid_signature",
        ),
        (lambda book: None, BATCH, 2, "stale_version"),
        (lambda book: None, "batch-B", 1, "batch_mismatch"),
        (
            lambda book: book["甲工厂"].__setitem__("I2", "foreign:1:token"),
            BATCH, 1, "foreign_token",
        ),
    ],
)
def test_security_rejections(change, batch: str, version: int, code: str) -> None:
    content, signature, lines = _source()
    with pytest.raises(IncomingWorkbookValidationError) as caught:
        _codec().parse(
            _edit(content, change), batch_id=batch, version=version,
            signature=signature, previous_lines=lines,
        )
    assert code in _codes(caught)


@pytest.mark.parametrize(
    ("limits", "code"),
    [
        (IncomingWorkbookLimits(max_source_bytes=1), "file_too_large"),
        (IncomingWorkbookLimits(max_zip_entries=1), "too_many_zip_entries"),
        (IncomingWorkbookLimits(max_uncompressed_bytes=1), "uncompressed_content_too_large"),
        (IncomingWorkbookLimits(max_worksheets=1), "too_many_worksheets"),
        (IncomingWorkbookLimits(max_data_rows=1), "too_many_data_rows"),
    ],
)
def test_five_hardening_limits_reject(limits: IncomingWorkbookLimits, code: str) -> None:
    content, signature, lines = _source()
    with pytest.raises(IncomingWorkbookValidationError) as caught:
        _codec(limits=limits).parse(
            content, batch_id=BATCH, version=1, signature=signature, previous_lines=lines
        )
    assert code in _codes(caught)
